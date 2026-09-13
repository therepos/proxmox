# Purpose: build real office files (xlsx, docx, pptx, pdf) on the share from text
# =============================================================================
# Loaded by server.py (write mode only):
#
#     from build_tools import register as register_build
#     register_build(mcp, _resolve, rel=_rel, guard=tool, link=download_link_for)
#
# Claude writes markdown or CSV; the file is assembled here and the result
# carries a signed download link, so the user gets a one-click download the
# same way they do in a chat with file output.
#
#   build_xlsx(path, csv | sheets)   CSV or markdown tables -> workbook
#   build_docx(path, markdown)       markdown -> Word document
#   build_pptx(path, markdown)       markdown -> slides ("# " = new slide)
#   build_pdf(path, markdown)        markdown -> PDF
#
# Markdown subset understood by the builders: # headings, paragraphs, - / 1.
# lists (two levels), | tables |, ``` code ```, **bold**, *italic*, `code`,
# and a line of --- for a page break (docx, pdf) or slide break (pptx).
#
# Deps (lazy): openpyxl, python-docx, python-pptx, reportlab.
# =============================================================================

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any, Callable

from mcp.types import ToolAnnotations

RW = ToolAnnotations(readOnlyHint=False, destructiveHint=False, idempotentHint=False, openWorldHint=False)
MAX_INPUT_CHARS = 2_000_000


def _need(module: str, pip_name: str):  # type: ignore[no-untyped-def]
    try:
        return __import__(module)
    except ImportError as e:
        raise RuntimeError(f"{pip_name} is not installed in the server environment. Re-run the installer and choose 'Update server'.") from e


# --- Markdown subset -> blocks ------------------------------------------------
_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*\n]+?\*(?!\*))")


def inlines(text: str) -> list[tuple[str, bool, bool, bool]]:
    """-> [(text, bold, italic, code)]"""
    out = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append((part[2:-2], True, False, False))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append((part[1:-1], False, False, True))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], False, True, False))
        else:
            out.append((part, False, False, False))
    return out


def plain(text: str) -> str:
    return "".join(t for t, *_ in inlines(text))


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


def parse_md(text: str) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    para: list[str] = []
    lst: dict[str, Any] | None = None
    table: list[list[str]] | None = None
    code: list[str] | None = None

    def flush() -> None:
        nonlocal para, lst, table
        if para:
            blocks.append({"type": "para", "text": " ".join(para)})
            para = []
        if lst:
            blocks.append(lst)
            lst = None
        if table:
            blocks.append({"type": "table", "rows": table})
            table = None

    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if code is not None:
            if line.strip().startswith("```"):
                blocks.append({"type": "code", "text": "\n".join(code)})
                code = None
            else:
                code.append(raw)
            continue
        if line.strip().startswith("```"):
            flush()
            code = []
            continue
        if not line.strip():
            flush()
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            blocks.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2).strip()})
            continue
        if re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", line):
            flush()
            blocks.append({"type": "break"})
            continue
        m = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if m:
            if para or table:
                flush()
            ordered = m.group(2)[0].isdigit()
            level = 1 if len(m.group(1)) >= 2 else 0
            if lst is None or lst["ordered"] != ordered and level == 0:
                if lst:
                    blocks.append(lst)
                lst = {"type": "list", "ordered": ordered, "items": []}
            lst["items"].append((level, m.group(3).strip()))
            continue
        if line.strip().startswith("|") and line.strip().endswith("|"):
            if para or lst:
                flush()
            cells = _split_row(line)
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c) and any(cells):
                continue  # header separator
            table = (table or []) + [cells]
            continue
        if lst is not None and line.startswith("  "):
            lvl, txt = lst["items"][-1]
            lst["items"][-1] = (lvl, txt + " " + line.strip())
            continue
        if lst or table:
            flush()
        para.append(line.strip())
    if code is not None:
        blocks.append({"type": "code", "text": "\n".join(code)})
    flush()
    return blocks


def _coerce(v: str) -> Any:
    s = v.strip()
    if s == "":
        return None
    if re.fullmatch(r"[-+]?0\d+", s):          # 007, -042: keep as text
        return s
    if re.fullmatch(r"[-+]?\d{1,15}", s):
        return int(s)
    if re.fullmatch(r"[-+]?(\d+\.\d*|\.\d+|\d+)([eE][-+]?\d+)?", s):
        try:
            return float(s)
        except ValueError:
            return s
    return s


def _rows_from_text(text: str) -> list[list[str]]:
    """CSV, or a markdown table, -> rows of strings."""
    stripped = text.strip()
    if stripped.startswith("|"):
        rows = []
        for line in stripped.splitlines():
            if not line.strip().startswith("|"):
                continue
            cells = _split_row(line)
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c) and any(cells):
                continue
            rows.append(cells)
        return rows
    return [r for r in csv.reader(io.StringIO(stripped)) if r]


# --- Registration ------------------------------------------------------------
def register(
    mcp,  # type: ignore[no-untyped-def]
    resolve: Callable[..., Path],
    *,
    rel: Callable[[Path], str],
    guard: Callable[[ToolAnnotations], Callable],
    link: Callable[[Path], dict[str, Any]],
) -> None:
    def _target(path: str, ext: str, overwrite: bool) -> Path:
        p = resolve(path, must_exist=False)
        if p.suffix.lower() != ext:
            p = p.with_name(p.name + ext)
        if p.exists() and not overwrite:
            raise FileExistsError(f"Destination exists: {rel(p)} (pass overwrite=true)")
        if p.is_dir():
            raise IsADirectoryError(f"Is a directory: {path}")
        p.parent.mkdir(parents=True, exist_ok=True)
        return p

    def _done(p: Path, **extra: Any) -> dict[str, Any]:
        out = {"saved": rel(p), "bytes": p.stat().st_size, **extra}
        out["download"] = link(p)
        out["next"] = "Give the download url to the user."
        return out

    def _check_len(*texts: str | None) -> None:
        if sum(len(t or "") for t in texts) > MAX_INPUT_CHARS:
            raise ValueError("Input too large; split the document.")

    # --- xlsx ----------------------------------------------------------------------
    @guard(RW)
    def build_xlsx(
        path: str,
        csv_text: str | None = None,
        sheets: list[dict[str, str]] | None = None,
        sheet_name: str = "Sheet1",
        header: bool = True,
        overwrite: bool = False,
    ) -> dict[str, Any]:
        """Create a real Excel workbook from CSV or markdown tables and return a download link.
        Numbers become numeric cells; the header row is bold, frozen and filterable; columns are
        sized to content.

        Args:
            path: destination .xlsx relative to the share root.
            csv_text: CSV (or a markdown table) for a single sheet.
            sheets: several sheets: [{"name": "Summary", "csv": "..."}, ...]. Overrides csv_text.
            sheet_name: name for the single sheet when csv_text is used.
            header: treat the first row as a header.
            overwrite: replace an existing file.
        """
        openpyxl = _need("openpyxl", "openpyxl")
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        specs = sheets or ([{"name": sheet_name, "csv": csv_text or ""}])
        _check_len(*(s.get("csv", "") for s in specs))
        if not any((s.get("csv") or "").strip() for s in specs):
            raise ValueError("No data: pass csv_text or sheets.")
        p = _target(path, ".xlsx", overwrite)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        used: set[str] = set()
        summary = []
        for i, spec in enumerate(specs, 1):
            name = re.sub(r"[\[\]:*?/\\]", "_", (spec.get("name") or f"Sheet{i}"))[:31] or f"Sheet{i}"
            base, n = name, 1
            while name.lower() in used:
                n += 1
                name = f"{base[:28]}_{n}"
            used.add(name.lower())
            ws = wb.create_sheet(name)
            rows = _rows_from_text(spec.get("csv") or "")
            widths: dict[int, int] = {}
            for r, row in enumerate(rows, 1):
                for c, val in enumerate(row, 1):
                    v: Any = val if (header and r == 1) else _coerce(val)
                    cell = ws.cell(row=r, column=c, value=v)
                    if header and r == 1:
                        cell.font = Font(bold=True)
                    widths[c] = max(widths.get(c, 0), min(len(str(val)), 60))
            for c, w in widths.items():
                ws.column_dimensions[get_column_letter(c)].width = max(8, w + 2)
            if header and rows:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
            summary.append({"sheet": name, "rows": len(rows), "cols": max((len(r) for r in rows), default=0)})
        wb.save(p)
        return _done(p, sheets=summary)

    # --- docx ----------------------------------------------------------------------
    @guard(RW)
    def build_docx(path: str, markdown: str, title: str | None = None, overwrite: bool = False) -> dict[str, Any]:
        """Create a real Word document from markdown and return a download link. Supports headings,
        paragraphs, bullet and numbered lists (two levels), tables, code blocks, **bold**, *italic*,
        `code`, and --- for a page break.

        Args:
            path: destination .docx relative to the share root.
            markdown: document body.
            title: optional document title placed above the body.
            overwrite: replace an existing file.
        """
        docx = _need("docx", "python-docx")
        from docx.shared import Pt

        _check_len(markdown)
        p = _target(path, ".docx", overwrite)
        doc = docx.Document()
        if title:
            doc.add_heading(title, 0)

        def add_runs(par, text: str) -> None:  # type: ignore[no-untyped-def]
            for t, b, i, c in inlines(text):
                run = par.add_run(t)
                run.bold = b or None
                run.italic = i or None
                if c:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9.5)

        n_tables = 0
        for blk in parse_md(markdown):
            t = blk["type"]
            if t == "heading":
                doc.add_heading(plain(blk["text"]), min(blk["level"], 9))
            elif t == "para":
                add_runs(doc.add_paragraph(), blk["text"])
            elif t == "list":
                for level, item in blk["items"]:
                    style = ("List Number" if blk["ordered"] else "List Bullet") + (" 2" if level else "")
                    try:
                        par = doc.add_paragraph(style=style)
                    except KeyError:
                        par = doc.add_paragraph(style="List Bullet")
                    add_runs(par, item)
            elif t == "table":
                rows = blk["rows"]
                width = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=width)
                tbl.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c in range(width):
                        cell = tbl.cell(r, c)
                        cell.text = ""
                        add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                n_tables += 1
                doc.add_paragraph()
            elif t == "code":
                par = doc.add_paragraph()
                run = par.add_run(blk["text"])
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif t == "break":
                doc.add_page_break()
        doc.save(p)
        return _done(p, paragraphs=len(doc.paragraphs), tables=n_tables)

    # --- pptx ----------------------------------------------------------------------
    @guard(RW)
    def build_pptx(path: str, markdown: str, overwrite: bool = False) -> dict[str, Any]:
        """Create a real PowerPoint deck from markdown and return a download link. Every "# Heading"
        (or a --- line) starts a new slide; the first slide is a title slide if the deck opens with a
        heading followed by a plain paragraph. Bullets, numbered lists, paragraphs and tables go on
        the slide body. A line starting with "Notes:" becomes speaker notes.

        Args:
            path: destination .pptx relative to the share root.
            markdown: deck content.
            overwrite: replace an existing file.
        """
        pptx = _need("pptx", "python-pptx")
        from pptx.util import Inches, Pt

        _check_len(markdown)
        p = _target(path, ".pptx", overwrite)
        prs = pptx.Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

        # split into slides
        slides: list[dict[str, Any]] = []
        cur: dict[str, Any] | None = None
        for blk in parse_md(markdown):
            if blk["type"] == "heading" and blk["level"] == 1 or blk["type"] == "break":
                title = plain(blk.get("text", ""))
                if cur is not None and not cur["title"] and not cur["blocks"] and not cur["notes"]:
                    cur["title"] = title          # break followed by heading: one slide, not two
                    continue
                cur = {"title": title, "blocks": [], "notes": []}
                slides.append(cur)
                continue
            if cur is None:
                cur = {"title": "", "blocks": [], "notes": []}
                slides.append(cur)
            if blk["type"] == "para" and blk["text"].lower().startswith("notes:"):
                cur["notes"].append(blk["text"][6:].strip())
            else:
                cur["blocks"].append(blk)
        if not slides:
            raise ValueError("No content.")

        first = slides[0]
        if first["title"] and len(first["blocks"]) == 1 and first["blocks"][0]["type"] == "para":
            s = prs.slides.add_slide(prs.slide_layouts[0])
            s.shapes.title.text = first["title"]
            s.placeholders[1].text = plain(first["blocks"][0]["text"])
            if first["notes"]:
                s.notes_slide.notes_text_frame.text = "\n".join(first["notes"])
            slides = slides[1:]

        for sd in slides:
            has_table = any(b["type"] == "table" for b in sd["blocks"])
            s = prs.slides.add_slide(prs.slide_layouts[5 if has_table else 1])
            s.shapes.title.text = sd["title"] or " "
            top = Inches(1.5)
            if not has_table:
                tf = s.placeholders[1].text_frame
                tf.clear()
                firstp = True
                for blk in sd["blocks"]:
                    lines: list[tuple[int, str]] = []
                    if blk["type"] == "list":
                        lines = blk["items"]
                    elif blk["type"] == "para":
                        lines = [(0, blk["text"])]
                    elif blk["type"] == "heading":
                        lines = [(0, "**" + blk["text"] + "**")]
                    elif blk["type"] == "code":
                        lines = [(0, "`" + ln + "`") for ln in blk["text"].splitlines()]
                    for level, text in lines:
                        par = tf.paragraphs[0] if firstp else tf.add_paragraph()
                        firstp = False
                        par.level = level
                        for t, b, i, c in inlines(text):
                            run = par.add_run()
                            run.text = t
                            run.font.bold = b or None
                            run.font.italic = i or None
                            if c:
                                run.font.name = "Consolas"
            else:
                for blk in sd["blocks"]:
                    if blk["type"] == "table":
                        rows = blk["rows"]
                        width = max(len(r) for r in rows)
                        height = Inches(0.4) * len(rows)
                        shape = s.shapes.add_table(len(rows), width, Inches(0.6), top, prs.slide_width - Inches(1.2), height)
                        for r, row in enumerate(rows):
                            for c in range(width):
                                cell = shape.table.cell(r, c)
                                cell.text = plain(row[c]) if c < len(row) else ""
                                for par in cell.text_frame.paragraphs:
                                    for run in par.runs:
                                        run.font.size = Pt(14)
                                        run.font.bold = True if r == 0 else None
                        top += height + Inches(0.3)
                    elif blk["type"] in ("para", "list"):
                        if blk["type"] == "para":
                            text = blk["text"]
                        else:
                            text = "\n".join((f"{n}. " if blk["ordered"] else "• ") + t for n, (_, t) in enumerate(blk["items"], 1))
                        box = s.shapes.add_textbox(Inches(0.6), top, prs.slide_width - Inches(1.2), Inches(1))
                        box.text_frame.word_wrap = True
                        box.text_frame.text = plain(text)
                        top += Inches(0.5) * (text.count("\n") + 1) + Inches(0.2)
            if sd["notes"]:
                s.notes_slide.notes_text_frame.text = "\n".join(sd["notes"])
        prs.save(p)
        return _done(p, slides=len(prs.slides))

    # --- pdf -----------------------------------------------------------------------
    @guard(RW)
    def build_pdf(path: str, markdown: str, title: str | None = None, overwrite: bool = False) -> dict[str, Any]:
        """Create a PDF from markdown and return a download link. Same markdown subset as build_docx
        (headings, lists, tables, code, bold/italic, --- page break). A4 portrait.

        Args:
            path: destination .pdf relative to the share root.
            markdown: document body.
            title: optional title placed above the body.
            overwrite: replace an existing file.
        """
        _need("reportlab", "reportlab")
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import (ListFlowable, ListItem, PageBreak, Paragraph, Preformatted,
                                        SimpleDocTemplate, Spacer, Table, TableStyle)

        _check_len(markdown)
        p = _target(path, ".pdf", overwrite)
        ss = getSampleStyleSheet()
        body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=10.5, leading=14, spaceAfter=6)
        code_style = ParagraphStyle("code", parent=ss["Code"], fontSize=8.5, leading=11, backColor=colors.whitesmoke, leftIndent=6)
        cell_style = ParagraphStyle("cell", parent=body, fontSize=9, leading=11, spaceAfter=0)

        def esc(s: str) -> str:
            return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        def rich(text: str) -> str:
            out = []
            for t, b, i, c in inlines(text):
                t = esc(t)
                if c:
                    t = f'<font face="Courier">{t}</font>'
                if b:
                    t = f"<b>{t}</b>"
                if i:
                    t = f"<i>{t}</i>"
                out.append(t)
            return "".join(out)

        story: list[Any] = []
        if title:
            story += [Paragraph(esc(title), ss["Title"]), Spacer(1, 6)]
        for blk in parse_md(markdown):
            t = blk["type"]
            if t == "heading":
                story.append(Paragraph(rich(blk["text"]), ss[f"Heading{min(blk['level'], 4)}"]))
            elif t == "para":
                story.append(Paragraph(rich(blk["text"]), body))
            elif t == "list":
                items = [ListItem(Paragraph(rich(txt), body), leftIndent=12 + 14 * lvl) for lvl, txt in blk["items"]]
                if blk["ordered"]:
                    story.append(ListFlowable(items, bulletType="1", leftIndent=14, bulletFontName="Helvetica", bulletFontSize=10))
                else:
                    story.append(ListFlowable(items, bulletType="bullet", start="\u2022", leftIndent=14, bulletFontName="Helvetica", bulletFontSize=10))
                story.append(Spacer(1, 4))
            elif t == "table":
                rows = blk["rows"]
                width = max(len(r) for r in rows)
                data = [[Paragraph(rich(r[c]) if c < len(r) else "", cell_style) for c in range(width)] for r in rows]
                tbl = Table(data, repeatRows=1, hAlign="LEFT", colWidths=[(A4[0] - 4 * cm) / width] * width)
                tbl.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ]))
                story += [tbl, Spacer(1, 8)]
            elif t == "code":
                story.append(Preformatted(blk["text"], code_style))
                story.append(Spacer(1, 6))
            elif t == "break":
                story.append(PageBreak())
        if not story:
            raise ValueError("No content.")
        doc = SimpleDocTemplate(str(p), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm, title=title or p.stem)
        doc.build(story)
        return _done(p)
