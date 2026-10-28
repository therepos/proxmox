"""
skill-awp.py - build an internal-audit work program (AWP) workbook in the house format from a spec,
or export the spec from an existing AWP in that format.

    python skill-awp.py build spec.json out.xlsx        # spec -> workbook
    python skill-awp.py template out.xlsx [--tests 5]   # blank template with placeholders
    python skill-awp.py export awp.xlsx spec.json       # existing AWP -> spec (same layout only)
    python skill-awp.py selftest                        # proves the environment works (builds the example and a template)

Workbook produced (series letter configurable, default "A"):
    Scope | A | A100 (RCM) | A200 (walkthrough index) | A201.. (walkthrough per sub-process)
          | A300 (testing index) | A301.. (one per test) | DA (data analytics)

Spec: see Part 3 of skill-awp.md. Every text field is optional - missing fields get placeholders in [brackets]
so the template is complete and reviewers can see what still has to be written.

Formatting (fixed - do not restyle by hand): Arial 8 throughout; yellow ref box with double border in row 1 of
the last column; Client / Process / Audit period in A3:B5 (formulas to the index sheet); grey D9D9D9 column
headers; DCE6F1 section rows in the RCM; 808080 bands with white bold text for sub-process titles and test
titles; B7DEE8 detail-column headers and 92D050 "Work Done" step headers on test sheets; bold red cross-
references (control numbers, WP refs); thin borders; wrapped text with row heights sized to the text;
gridlines off; tab colours black (Scope) and 92D050 (all others).
"""
import json, math, re, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string

# ----------------------------------------------------------------------------- styles
FN = 'Arial'; SZ = 8
def F(bold=False, color=None): return Font(name=FN, size=SZ, bold=bold, color=color)
RED = 'FFFF0000'; WHITE = 'FFFFFFFF'
FILL = lambda rgb: PatternFill('solid', start_color=rgb)
F_HDR = FILL('D9D9D9'); F_SECTION = FILL('DCE6F1'); F_BAND = FILL('808080'); F_DETAIL = FILL('B7DEE8'); F_STEPS = FILL('92D050'); F_REF = FILL('FFFF00'); F_BLACK = FILL('000000'); F_ACCENT = FILL('4BACC6')
THIN = Side(style='thin'); DOUBLE = Side(style='double')
B_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN); B_DOUBLE = Border(left=DOUBLE, right=DOUBLE, top=DOUBLE, bottom=DOUBLE); B_BOTTOM = Border(bottom=THIN)
TOP = Alignment(vertical='top'); TOPW = Alignment(vertical='top', wrap_text=True); LEFTW = Alignment(horizontal='left', vertical='top', wrap_text=True)
CTR = Alignment(horizontal='center', vertical='top'); CTRW = Alignment(horizontal='center', vertical='top', wrap_text=True); CTRM = Alignment(horizontal='center', vertical='center', wrap_text=True)
TAB_GREEN = '92D050'; TAB_BLACK = '000000'
LINE_PT = 11.25   # points per wrapped line at Arial 8

def ph(v, text):
    """Placeholder when a spec value is missing."""
    return v if (v is not None and str(v).strip() != '') else f'[{text}]'

def est_lines(text, width_chars):
    if text is None: return 1
    s = str(text); n = 0
    for para in s.split('\n'):
        n += max(1, math.ceil(len(para) / max(width_chars * 1.15, 1)))
    return n

def fit_height(ws, row, cells_widths, min_h=24.75):
    """Set the row height from the longest wrapped text among (value, column width) pairs."""
    lines = max([est_lines(v, w) for v, w in cells_widths] + [1])
    ws.row_dimensions[row].height = max(min_h, min(409, lines * LINE_PT + 3))

def put(ws, ref, value, font=None, fill=None, align=None, border=None):
    c = ws[ref]; c.value = value; c.font = font or F()
    if fill: c.fill = fill
    if align: c.alignment = align
    if border: c.border = border
    return c

def ref_box(ws, col_letter, ref):
    put(ws, f'{col_letter}1', ref, F(True), F_REF, CTRM, B_DOUBLE); ws.row_dimensions[1].height = 25.5

def header(ws, spec, index_sheet, labels=('Client', 'Process', 'Audit period '), signoff=True, signoff_col=None):
    """A3:B5 client / process / period. On the index sheet the values are typed; elsewhere they are formulas."""
    for i, lab in enumerate(labels):
        put(ws, f'A{3+i}', lab, F(True))
        if ws.title == index_sheet:
            put(ws, f'B{3+i}', [spec.get('client', ''), spec.get('process', ''), spec.get('period', '')][i] or f'[{lab.strip()}]', F(), align=TOP)
        else:
            put(ws, f'B{3+i}', f"='{index_sheet}'!B{3+i}", F(), align=TOP)
    if signoff and signoff_col:
        c1 = signoff_col; c2 = get_column_letter(column_index_from_string(c1) + 2)
        put(ws, f'{c1}3', 'Prepared by:', F(True)); put(ws, f'{c2}3', 'Date:', F(True))
        put(ws, f'{c1}5', 'Reviewed by:', F(True)); put(ws, f'{c2}5', 'Date:', F(True))
        put(ws, f'{get_column_letter(column_index_from_string(c1)+1)}3', spec.get('prepared_by', ''), F())
        put(ws, f'{get_column_letter(column_index_from_string(c1)+1)}5', spec.get('reviewed_by', ''), F())

def widths(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v

def finish(ws, tab=TAB_GREEN):
    ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = tab

# ----------------------------------------------------------------------------- sheets
def build_scope(wb, spec, P):
    ws = wb.create_sheet('Scope')
    for col, h in zip('ABCD', ['No', 'Key Process', 'Indicative area of focus', 'Covered?']):
        put(ws, f'{col}1', h, F(True, WHITE if col != 'D' else None), F_BLACK if col != 'D' else F_ACCENT, CTRW, B_THIN)
    r = 2
    for i, kp in enumerate(spec.get('scope') or [{'key_process': None, 'areas': [None, None], 'covered': 'Yes'}], 1):
        areas = kp.get('areas') or [None]
        start = r
        for a in areas:
            put(ws, f'A{r}', kp.get('no', i), F(), align=CTRW, border=B_THIN)
            put(ws, f'B{r}', ph(kp.get('key_process'), 'Key process'), F(), align=LEFTW, border=B_THIN)
            put(ws, f'C{r}', ph(a, 'Indicative area of focus - one line per area'), F(), align=LEFTW, border=B_THIN)
            put(ws, f'D{r}', kp.get('covered', 'Yes'), F(), align=CTR, border=B_THIN)
            fit_height(ws, r, [(ws[f'C{r}'].value, 80)])
            r += 1
        if len(areas) > 1:
            ws.merge_cells(f'A{start}:A{r-1}'); ws.merge_cells(f'B{start}:B{r-1}')
    widths(ws, {'A': 6, 'B': 49.1, 'C': 80, 'D': 42.4}); finish(ws, TAB_BLACK)

def build_index(wb, spec, P, name, title, items, ref_col, w):
    """Generic two-column index sheet (A, A200, A300): number | title | ref."""
    ws = wb.create_sheet(name); ref_box(ws, ref_col, name)
    header(ws, spec, P, signoff=spec.get('signoff', True) and name == P, signoff_col='D' if name == P else None)
    put(ws, 'B7', title, F(True), align=TOPW, border=B_BOTTOM); put(ws, f'{ref_col}7', 'Ref', F(True), align=CTRW, border=B_BOTTOM)
    r = 9
    for i, (txt, ref) in enumerate(items, 1):
        put(ws, f'A{r}', i, F(), align=CTR); put(ws, f'B{r}', txt, F(), align=LEFTW); put(ws, f'{ref_col}{r}', ref, F(True, RED), align=CTRM)
        fit_height(ws, r, [(txt, w['B'])], min_h=13.5); r += 2
    widths(ws, w); finish(ws)
    return ws

def build_rcm(wb, spec, P):
    name = f'{P}100'; ws = wb.create_sheet(name); ref_box(ws, 'H', name)
    header(ws, spec, P, labels=('Client:', 'Process:', 'Review period:'))
    hdrs = ['S/N', 'Sub-process', 'Risk', 'Control No.', 'Finding for Design Effectiveness', 'Audit Strategy', 'Testing Result', 'WP Ref']
    for j, h in enumerate(hdrs):
        put(ws, f'{get_column_letter(j+1)}7', h, F(True), F_HDR, Alignment(horizontal='left' if j == 1 else 'center', vertical='center', wrap_text=True), B_THIN)
    ws.row_dimensions[7].height = 30.75
    r = 8; sn = 1
    sections = spec.get('rcm', {}).get('sections') or [{'title': None, 'rows': [{}, {}]}]
    for k, sec in enumerate(sections, 1):
        put(ws, f'A{r}', ph(sec.get('title'), f'{k}. Section title'), F(True), F_SECTION, TOP, B_THIN)
        for j in range(2, 9): ws[f'{get_column_letter(j)}{r}'].border = B_THIN; ws[f'{get_column_letter(j)}{r}'].fill = F_SECTION
        ws.merge_cells(f'A{r}:H{r}'); ws.row_dimensions[r].height = 22.5; r += 1
        for row in sec.get('rows') or [{}]:
            vals = [row.get('sn', sn), ph(row.get('sub_process'), 'Sub-process'), ph(row.get('risk'), 'Risk - what could go wrong and the consequence'), ph(row.get('control_no'), f'{P}2xx-n'),
                    row.get('de_finding') or 'To be concluded after the walkthrough.', ph(row.get('audit_strategy'), 'Test n: title\nCheck that:\n1. ...'), row.get('testing_result') or '', ph(row.get('wp_ref'), f'{P}3xx')]
            aligns = [CTR, LEFTW, LEFTW, CTRW, LEFTW, LEFTW, LEFTW, CTRW]
            for j, (v, al) in enumerate(zip(vals, aligns)):
                put(ws, f'{get_column_letter(j+1)}{r}', v, F(True, RED) if j in (3, 7) else F(), align=al, border=B_THIN)
            fit_height(ws, r, [(vals[2], 69.6), (vals[4], 67.3), (vals[5], 83.1), (vals[6], 70.9), (vals[1], 44.9)])
            r += 1; sn += 1
    widths(ws, {'A': 14.4, 'B': 44.9, 'C': 69.6, 'D': 15.9, 'E': 67.3, 'F': 83.1, 'G': 70.9, 'H': 13.6}); ws.freeze_panes = 'C8'; finish(ws)

def build_walkthrough(wb, spec, P, wt):
    name = wt['ref']; ws = wb.create_sheet(name); ref_box(ws, 'K', name)
    header(ws, spec, P)
    put(ws, 'A7', f"Subprocess title: {ph(wt.get('title'), 'Sub-process title')}", F(True, WHITE), F_BAND, TOP, B_THIN)
    for col in 'BCDEFGHIJK': ws[f'{col}7'].fill = F_BAND; ws[f'{col}7'].border = B_THIN
    ws.merge_cells('A7:K7')
    bg = 'Background:\n\n' + ph(wt.get('background'), 'Describe the process, the systems used, the people involved and any change during the period that matters for this sub-process.')
    put(ws, 'A8', bg, F(), align=LEFTW, border=B_THIN); ws.merge_cells('A8:K8'); fit_height(ws, 8, [(bg, 260)], min_h=40)
    hdrs = ['No.', 'Activity', 'Performed By', 'Frequency', 'Control Reference No.', 'Existing controls / mitigating measures', 'Control Owner', 'Control Classification', 'Control Type', 'Finding for design effectiveness', 'Walkthrough Document']
    for j, h in enumerate(hdrs):
        put(ws, f'{get_column_letter(j+1)}9', h, F(True), F_HDR, TOPW if j < 2 else CTRW, B_THIN)
    ws.row_dimensions[9].height = 22.5
    r = 10; n = 1; cref = 1
    for s, sec in enumerate(wt.get('sections') or [{'title': None, 'activities': [{}]}], 1):
        put(ws, f'A{r}', ph(sec.get('title'), f'{s}.{1} Sub-section title'), F(True), align=LEFTW, border=B_THIN)
        for col in 'BCDEFGHIJK': ws[f'{col}{r}'].border = B_THIN
        ws.merge_cells(f'A{r}:K{r}'); ws.row_dimensions[r].height = 13.5; r += 1
        for act in sec.get('activities') or [{}]:
            if act.get('refer_to'):
                vals = [act.get('no', n), act['refer_to']] + [''] * 9
            else:
                cr = act.get('control_ref')
                if cr is None and (act.get('control') or act.get('owner')): cr = f'{name}-{cref}'; cref += 1
                vals = [act.get('no', n), ph(act.get('activity'), 'Activity - who does what, when, in which system'), ph(act.get('performed_by'), 'Performed by'), ph(act.get('frequency'), 'Frequency'),
                        cr or '', ph(act.get('control'), 'Control - what prevents or detects the risk, and the evidence'), ph(act.get('owner'), 'Control owner'), ph(act.get('classification'), 'Preventive / Detective'), ph(act.get('type'), 'Manual / IT'),
                        act.get('de_finding') or '', ph(act.get('document'), 'Document to sight')]
            aligns = [CTR, TOPW, CTRW, CTRW, CTR, LEFTW, CTRW, CTRW, CTRW, LEFTW, CTRW]
            for j, (v, al) in enumerate(zip(vals, aligns)):
                put(ws, f'{get_column_letter(j+1)}{r}', v, F(True, RED) if j == 4 else F(), align=al, border=B_THIN)
            fit_height(ws, r, [(vals[1], 42.6), (vals[5], 50.4), (vals[9], 59.1), (vals[2], 19), (vals[10], 16)], min_h=24.75)
            r += 1; n += 1
    r += 1
    qtitle = wt.get('questions_title') or 'Walkthrough questions for the team lead (to be used during the walkthrough session)'
    put(ws, f'A{r}', qtitle, F(True, WHITE), F_BAND, TOP, B_THIN)
    for col in 'BCDEFGHIJK': ws[f'{col}{r}'].fill = F_BAND; ws[f'{col}{r}'].border = B_THIN
    ws.merge_cells(f'A{r}:K{r}'); r += 1
    qh = [('A', 'No.'), ('B', 'Question'), ('E', 'Why we are asking, including the prior-year context'), ('G', 'Documents to request during the session'), ('I', 'Proposed interviewee'), ('J', 'Response noted during the walkthrough'), ('K', 'Control Reference No.')]
    for col, h in qh: put(ws, f'{col}{r}', h, F(True), F_HDR, CTRW, B_THIN)
    for col in 'ABCDEFGHIJK': ws[f'{col}{r}'].border = B_THIN; ws[f'{col}{r}'].fill = F_HDR
    ws.merge_cells(f'B{r}:D{r}'); ws.merge_cells(f'E{r}:F{r}'); ws.merge_cells(f'G{r}:H{r}'); ws.row_dimensions[r].height = 22.5; r += 1
    for i, q in enumerate(wt.get('questions') or [{}, {}, {}], 1):
        vals = {'A': q.get('no', i), 'B': ph(q.get('question'), 'Open question that makes the interviewee describe the step'), 'E': ph(q.get('why'), 'Why we ask - the risk, the prior-year point or the change we want to understand'), 'G': ph(q.get('documents'), 'Documents to request'), 'I': ph(q.get('interviewee'), 'Interviewee'), 'J': q.get('response') or '', 'K': q.get('control_ref') or ''}
        for col, v in vals.items():
            put(ws, f'{col}{r}', v, F(True, RED) if col == 'K' else F(), align=CTR if col == 'A' else (CTRW if col in 'IK' else LEFTW), border=B_THIN)
        for col in 'ABCDEFGHIJK': ws[f'{col}{r}'].border = B_THIN
        ws.merge_cells(f'B{r}:D{r}'); ws.merge_cells(f'E{r}:F{r}'); ws.merge_cells(f'G{r}:H{r}')
        fit_height(ws, r, [(vals['B'], 80), (vals['E'], 66), (vals['G'], 32), (vals['J'], 59)], min_h=24.75); r += 1
    widths(ws, {'A': 12.6, 'B': 42.6, 'C': 19, 'D': 19.1, 'E': 16, 'F': 50.4, 'G': 16.1, 'H': 16.6, 'I': 12.9, 'J': 59.1, 'K': 16.1}); ws.freeze_panes = 'A10'; finish(ws)

def build_test(wb, spec, P, t):
    name = t['ref']; ws = wb.create_sheet(name)
    detail = t.get('detail_columns') or ['Document / item', 'Reference', 'Date', 'Amount (S$)']
    n_steps = int(t.get('n_steps') or max(1, len(re.findall(r'^\s*\d+\.', str(t.get('steps') or ''), re.M)) or 6))
    n_rows = int(t.get('n_rows') or 10)
    first_step = 3 + len(detail)                       # B = S/N, C.. = detail
    last = first_step + n_steps                        # remarks column
    L = get_column_letter; last_l = L(last)
    ref_box(ws, last_l, name); ws.row_dimensions[1].height = 14.25
    header(ws, spec, P)
    put(ws, 'A7', f"Test {t.get('no', '')}".strip(), F(True, WHITE), F_BAND, TOP, B_THIN); put(ws, 'B7', ph(t.get('title'), 'Test title'), F(True, WHITE), F_BAND, TOP)
    for j in range(2, last + 1): ws[f'{L(j)}7'].fill = F_BAND
    merge_to = L(min(last - 1, 9))
    fields = [('Objective: ', ph(t.get('objective'), 'To check that ...')), ('Basis of Selection:', ph(t.get('basis'), 'Haphazard / judgemental, with emphasis on ... / Not applicable')), ('Population:', ph(t.get('population'), 'Population and its source listing')),
              ('Sample Size:', ph(t.get('sample_size'), 'n, or all items if the population is smaller')), ('Source:', ph(t.get('source'), 'Documents and listings used')), ('Audit Steps:', ph(t.get('steps'), f"Test {t.get('no', 'n')}: title\nCheck that:\n1. ...\n2. ..."))]
    r = 9
    for lab, val in fields:
        put(ws, f'A{r}', lab, F(True), align=TOPW); put(ws, f'B{r}', val, F(), align=LEFTW); ws.merge_cells(f'B{r}:{merge_to}{r}')
        fit_height(ws, r, [(val, 150)], min_h=14.25); r += 2
    # work-done table
    put(ws, f'A{r}', 'Work done:', F(True), align=TOP)
    put(ws, f'B{r}', ph(t.get('detail_title'), 'Sample details'), F(True), F_DETAIL, CTR, B_THIN)
    for j in range(3, first_step): ws[f'{L(j)}{r}'].fill = F_DETAIL; ws[f'{L(j)}{r}'].border = B_THIN
    ws.merge_cells(f'B{r}:{L(first_step-1)}{r}')
    put(ws, f'{L(first_step)}{r}', 'Work Done', F(True), F_STEPS, CTRW, B_THIN)
    for j in range(first_step + 1, last): ws[f'{L(j)}{r}'].fill = F_STEPS; ws[f'{L(j)}{r}'].border = B_THIN
    if n_steps > 1: ws.merge_cells(f'{L(first_step)}{r}:{L(last-1)}{r}')
    put(ws, f'{last_l}{r}', 'Remarks', F(True), F_DETAIL, CTR, B_THIN); ws.row_dimensions[r].height = 15
    r += 1
    put(ws, f'B{r}', 'S/N', F(True), F_DETAIL, CTRW, B_THIN)
    for j, h in enumerate(detail): put(ws, f'{L(3+j)}{r}', h, F(True), F_DETAIL, CTRW, B_THIN)
    for s in range(n_steps): put(ws, f'{L(first_step+s)}{r}', s + 1, F(True), F_STEPS, CTRW, B_THIN)
    put(ws, f'{last_l}{r}', 'Exceptions and remarks', F(True), F_DETAIL, CTRW, B_THIN); ws.row_dimensions[r].height = 45
    hdr_row = r; r += 1
    rows = t.get('rows') or [{} for _ in range(n_rows)]
    for i, row in enumerate(rows, 1):
        put(ws, f'B{r}', i, F(), align=CTR, border=B_THIN)
        vals = row.get('values') or []
        formulas = t.get('detail_formulas') or {}
        for j in range(len(detail)):
            v = vals[j] if j < len(vals) else None
            if v is None and detail[j] in formulas: v = formulas[detail[j]].replace('{r}', str(r))
            put(ws, f'{L(3+j)}{r}', v, F(), align=TOPW, border=B_THIN)
        steps = row.get('steps') or []
        for s in range(n_steps):
            put(ws, f'{L(first_step+s)}{r}', steps[s] if s < len(steps) else None, F(), align=CTRW, border=B_THIN)
        put(ws, f'{last_l}{r}', row.get('remarks'), F(), align=LEFTW, border=B_THIN)
        ws.row_dimensions[r].height = 33.75; r += 1
    r += 1
    for lab, key, default in [('Notes:', 'notes', 'Nil'), ('Findings:', 'findings', 'Nil'), ('Conclusion:', 'conclusion', 'To be concluded on completion of fieldwork.')]:
        put(ws, f'A{r}', lab, F(True), align=LEFTW); put(ws, f'B{r}', t.get(key) or default, F(), align=TOP); r += 2
    w = {'A': 20.6, 'B': 7.6, last_l: 47.4}
    for j, h in enumerate(detail): w[L(3+j)] = max(14, min(42, int(len(h) * 1.4) + 8))
    for s in range(n_steps): w[L(first_step+s)] = 8.9
    widths(ws, w); ws.freeze_panes = f'C{hdr_row+1}'; finish(ws)

def build_da(wb, spec, P):
    ws = wb.create_sheet('DA'); ref_box(ws, 'G', 'DA'); header(ws, spec, P)
    put(ws, 'B7', 'Data analytics', F(True), align=TOPW); put(ws, 'G7', 'Ref ', F(True), align=CTRW)
    hdrs = ['S/N', 'Analytic', 'Data source', 'What the analytic is designed to identify', 'Feeds', 'Result', 'Status']
    for j, h in enumerate(hdrs): put(ws, f'{get_column_letter(j+1)}9', h, F(True), F_HDR, CTRW, B_THIN)
    ws.row_dimensions[9].height = 22.5
    r = 10
    for i, a in enumerate(spec.get('da', {}).get('analytics') or [{}, {}, {}], 1):
        vals = [a.get('sn', i), ph(a.get('analytic'), 'Analytic - what is compared with what'), ph(a.get('source'), 'Listings required'), ph(a.get('identifies'), 'What an exception looks like'), ph(a.get('feeds'), f'{P}3xx'), a.get('result') or '', a.get('status') or 'Planned']
        aligns = [CTR, LEFTW, LEFTW, LEFTW, CTRW, LEFTW, CTRW]
        for j, (v, al) in enumerate(zip(vals, aligns)): put(ws, f'{get_column_letter(j+1)}{r}', v, F(j == 4), align=al, border=B_THIN)
        fit_height(ws, r, [(vals[1], 64), (vals[3], 60), (vals[2], 18.9)]); r += 1
    r += 1
    note = spec.get('da', {}).get('note') or 'The analytics are to be run on the population listings requested in advance of fieldwork, so that the results can be used to select the samples for the tests above before the start of fieldwork.'
    put(ws, f'A{r}', note, F(), align=LEFTW); ws.merge_cells(f'A{r}:G{r}'); fit_height(ws, r, [(note, 200)], min_h=24)
    widths(ws, {'A': 13.6, 'B': 64, 'C': 18.9, 'D': 60, 'E': 9.7, 'F': 30, 'G': 11}); finish(ws)

# ----------------------------------------------------------------------------- orchestration
def normalise(spec):
    P = spec.get('prefix') or 'A'
    wts = spec.get('walkthroughs') or [{'title': None}, {'title': None}]
    for i, wt in enumerate(wts, 1): wt.setdefault('ref', f'{P}{200+i}')
    tests = spec.get('tests') or [{'title': None} for _ in range(3)]
    for i, t in enumerate(tests, 1): t.setdefault('ref', f'{P}{300+i}'); t.setdefault('no', i)
    spec['walkthroughs'] = wts; spec['tests'] = tests; spec['prefix'] = P
    return spec

def build_awp(spec, out_path):
    spec = normalise(json.loads(json.dumps(spec))); P = spec['prefix']
    wb = Workbook(); wb.remove(wb.active)
    build_scope(wb, spec, P)
    build_index(wb, spec, P, P, 'Index', [('Risk and control matrix', f'{P}100'), ('Process walkthrough', f'{P}200'), ('Testing of controls', f'{P}300'), ('Data Analytics', 'DA')], 'G', {'A': 12, 'B': 64, 'C': 8.6, 'D': 12.1, 'E': 7.4, 'F': 8.4, 'G': 14.6})
    build_rcm(wb, spec, P)
    build_index(wb, spec, P, f'{P}200', 'Process walkthrough', [(wt.get('index_text') or ph(wt.get('title'), 'Sub-process'), wt['ref']) for wt in spec['walkthroughs']], 'G', {'A': 13.6, 'B': 64, 'C': 18.9, 'D': 7, 'E': 8.4, 'F': 7.4, 'G': 11})
    for wt in spec['walkthroughs']: build_walkthrough(wb, spec, P, wt)
    build_index(wb, spec, P, f'{P}300', 'Testing', [(ph(t.get('title'), 'Test title'), t['ref']) for t in spec['tests']], 'E', {'A': 13.6, 'B': 64, 'C': 18.9, 'D': 7, 'E': 9.7, 'F': 11})
    for t in spec['tests']: build_test(wb, spec, P, t)
    build_da(wb, spec, P)
    wb.save(out_path); return out_path

def template_spec(n_walkthroughs=2, n_tests=3, n_analytics=3, prefix='A'):
    return {'prefix': prefix, 'client': '', 'process': '', 'period': '', 'signoff': True,
            'scope': [{'no': 1, 'key_process': '', 'areas': ['', ''], 'covered': 'Yes'}, {'no': 2, 'key_process': '', 'areas': ['', ''], 'covered': 'Yes'}],
            'rcm': {'sections': [{'title': '', 'rows': [{}, {}]}, {'title': '', 'rows': [{}]}]},
            'walkthroughs': [{'title': '', 'background': '', 'sections': [{'title': '', 'activities': [{}, {}]}], 'questions': [{}, {}, {}]} for _ in range(n_walkthroughs)],
            'tests': [{'title': '', 'n_steps': 6, 'n_rows': 10} for _ in range(n_tests)],
            'da': {'analytics': [{} for _ in range(n_analytics)]}}

# ----------------------------------------------------------------------------- export (reverse)
def export_spec(path):
    """Read an AWP built in this layout and return the spec (used to turn a finished AWP into a reusable example)."""
    wb = load_workbook(path)
    idx = [s for s in wb.sheetnames if re.fullmatch(r'[A-Z]', s)]
    P = idx[0] if idx else 'A'
    A = wb[P]
    spec = {'prefix': P, 'client': A['B3'].value, 'process': A['B4'].value, 'period': A['B5'].value, 'scope': [], 'rcm': {'sections': []}, 'walkthroughs': [], 'tests': [], 'da': {'analytics': []}}
    if 'Scope' in wb.sheetnames:
        ws = wb['Scope']; cur = None
        for r in range(2, ws.max_row + 1):
            if ws[f'C{r}'].value is None: continue
            if ws[f'B{r}'].value: cur = {'no': ws[f'A{r}'].value, 'key_process': ws[f'B{r}'].value, 'areas': [], 'covered': ws[f'D{r}'].value}; spec['scope'].append(cur)
            if cur: cur['areas'].append(ws[f'C{r}'].value)
    ws = wb[f'{P}100']; sec = None
    for r in range(8, ws.max_row + 1):
        a = ws[f'A{r}'].value
        if a is None: continue
        if ws[f'C{r}'].value is None and ws[f'B{r}'].value is None:
            sec = {'title': a, 'rows': []}; spec['rcm']['sections'].append(sec); continue
        if sec is None: sec = {'title': '', 'rows': []}; spec['rcm']['sections'].append(sec)
        sec['rows'].append({'sn': a, 'sub_process': ws[f'B{r}'].value, 'risk': ws[f'C{r}'].value, 'control_no': ws[f'D{r}'].value, 'de_finding': ws[f'E{r}'].value, 'audit_strategy': ws[f'F{r}'].value, 'testing_result': ws[f'G{r}'].value, 'wp_ref': ws[f'H{r}'].value})
    for name in wb.sheetnames:
        if re.fullmatch(P + r'2\d\d', name) and name != f'{P}200':
            ws = wb[name]; wt = {'ref': name, 'title': str(ws['A7'].value or '').replace('Subprocess title: ', ''), 'background': str(ws['A8'].value or '').replace('Background:\n\n', ''), 'sections': [], 'questions': []}
            idxws = wb[f'{P}200']
            for r in range(9, idxws.max_row + 1):
                if idxws[f'G{r}'].value == name: wt['index_text'] = idxws[f'B{r}'].value
            r = 10; sec = None; mode = 'act'
            while r <= ws.max_row:
                a, b = ws[f'A{r}'].value, ws[f'B{r}'].value
                if a and str(a).startswith('Walkthrough questions'): mode = 'q'; wt['questions_title'] = a; r += 2; continue
                if mode == 'act':
                    if a is not None and b is None and not isinstance(a, (int, float)): sec = {'title': a, 'activities': []}; wt['sections'].append(sec)
                    elif a is not None and b is not None:
                        if sec is None: sec = {'title': '', 'activities': []}; wt['sections'].append(sec)
                        act = {'no': a, 'activity': b, 'performed_by': ws[f'C{r}'].value, 'frequency': ws[f'D{r}'].value, 'control_ref': ws[f'E{r}'].value, 'control': ws[f'F{r}'].value, 'owner': ws[f'G{r}'].value, 'classification': ws[f'H{r}'].value, 'type': ws[f'I{r}'].value, 'de_finding': ws[f'J{r}'].value, 'document': ws[f'K{r}'].value}
                        if all(act[k] in (None, '') for k in ('performed_by', 'control', 'owner')): act = {'no': a, 'refer_to': b}
                        sec['activities'].append(act)
                else:
                    if a is not None and b is not None and str(b) != 'Question':
                        wt['questions'].append({'no': a, 'question': b, 'why': ws[f'E{r}'].value, 'documents': ws[f'G{r}'].value, 'interviewee': ws[f'I{r}'].value, 'response': ws[f'J{r}'].value, 'control_ref': ws[f'K{r}'].value})
                r += 1
            spec['walkthroughs'].append(wt)
    for name in wb.sheetnames:
        if re.fullmatch(P + r'3\d\d', name) and name != f'{P}300':
            ws = wb[name]; t = {'ref': name, 'no': int(re.sub(r'\D', '', str(ws['A7'].value)) or 0), 'title': ws['B7'].value}
            lab = {'Objective': 'objective', 'Basis of Selection': 'basis', 'Population': 'population', 'Sample Size': 'sample_size', 'Source': 'source', 'Audit Steps': 'steps', 'Notes': 'notes', 'Findings': 'findings', 'Conclusion': 'conclusion'}
            wd = None
            for r in range(9, ws.max_row + 1):
                a = ws[f'A{r}'].value
                if a is None: continue
                key = str(a).strip().rstrip(':').strip()
                if key in lab: t[lab[key]] = ws[f'B{r}'].value
                if key == 'Work done': wd = r
            if wd:
                t['detail_title'] = ws[f'B{wd}'].value
                hdr = wd + 1; cols = []; steps = 0; last = None
                for j in range(3, ws.max_column + 1):
                    v = ws.cell(row=hdr, column=j).value
                    if v is None: continue
                    if isinstance(v, (int, float)) or (isinstance(v, str) and v.isdigit()): steps += 1
                    elif str(v).startswith('Exceptions'): last = j
                    else: cols.append(v)
                t['detail_columns'] = cols; t['n_steps'] = steps
                n = 0; r = hdr + 1
                while ws[f'B{r}'].value is not None: n += 1; r += 1
                t['n_rows'] = n
                first = hdr + 1; formulas = {}
                for j, h in enumerate(cols):
                    v = ws.cell(row=first, column=3 + j).value
                    if isinstance(v, str) and v.startswith('='):
                        formulas[h] = re.sub(r'([A-Z]{1,2})' + str(first) + r'(?!\d)', r'\1{r}', v)
                if formulas: t['detail_formulas'] = formulas
            spec['tests'].append(t)
    if 'DA' in wb.sheetnames:
        ws = wb['DA']
        for r in range(10, ws.max_row + 1):
            if ws[f'B{r}'].value and ws[f'A{r}'].value is not None:
                spec['da']['analytics'].append({'sn': ws[f'A{r}'].value, 'analytic': ws[f'B{r}'].value, 'source': ws[f'C{r}'].value, 'identifies': ws[f'D{r}'].value, 'feeds': ws[f'E{r}'].value, 'result': ws[f'F{r}'].value, 'status': ws[f'G{r}'].value})
            elif ws[f'A{r}'].value and ws[f'B{r}'].value is None and len(str(ws[f'A{r}'].value)) > 40: spec['da']['note'] = ws[f'A{r}'].value
    return spec

def selftest():
    import os, tempfile
    here = os.path.dirname(os.path.abspath(__file__))
    ex = os.path.join(here, '..', 'examples', 'example_spec_commercial_leasing.json')
    out = tempfile.mkdtemp()
    build_awp(template_spec(n_tests=12), os.path.join(out, 'template.xlsx'))
    spec = json.load(open(ex, encoding='utf-8')) if os.path.exists(ex) else template_spec()
    path = build_awp(spec, os.path.join(out, 'example.xlsx'))
    back = export_spec(path)
    wb = load_workbook(path)
    ok = len(wb.sheetnames) == 4 + 2 + len(spec.get('walkthroughs', [])) + len(spec.get('tests', [])) and len(back['tests']) == len(spec.get('tests', []))
    print('SELFTEST', 'PASSED' if ok else 'FAILED', '-', len(wb.sheetnames), 'sheets built in', out)
    return ok

if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else ''
    if cmd == 'selftest': sys.exit(0 if selftest() else 1)
    if cmd == 'build': build_awp(json.load(open(sys.argv[2])), sys.argv[3]); print('built', sys.argv[3])
    elif cmd == 'template':
        n = int(sys.argv[sys.argv.index('--tests') + 1]) if '--tests' in sys.argv else 3
        build_awp(template_spec(n_tests=n), sys.argv[2]); print('template', sys.argv[2])
    elif cmd == 'export': json.dump(export_spec(sys.argv[2]), open(sys.argv[3], 'w'), indent=2, ensure_ascii=False, default=str); print('exported', sys.argv[3])
    else: print(__doc__)
