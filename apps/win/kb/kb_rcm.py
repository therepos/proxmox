# Purpose: Stage 2. Parse RCM / audit programme tables in kb.db into rows (risk, control, test procedure) -> table rcm_rows and rcm_rows.xlsx.
#   python kb_rcm.py [--db kb.db] [--types "RCM,Audit programme"] [--include-hidden]   (--include-hidden also parses RCM-like tables in other doc types)
import sqlite3, argparse, re
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ap = argparse.ArgumentParser(); ap.add_argument('--db', default='kb.db'); ap.add_argument('--types', default='RCM,Audit programme')
ap.add_argument('--include-hidden', action='store_true'); ap.add_argument('--out', default='rcm_rows.xlsx'); a = ap.parse_args()
db = sqlite3.connect(a.db); db.row_factory = sqlite3.Row

# column role detection from a header row. Order matters: first matching role wins per cell, first cell wins per role.
ROLES = [
    ('test',       r'test|procedure|audit step|work ?step|audit program|verification|steps? to|approach'),
    ('control',    r'control activit|control descr|key control|controls?\b(?!.*(ref|no\.?|owner|type|freq|rating|id))|mitigat'),
    ('risk',       r'risk descr|risks?\b(?!.*(rating|ref|no\.?|id|level|score|category|type|owner))|what could go wrong|wcgw'),
    ('objective',  r'objective|assertion'),
    ('subprocess', r'sub.?process|activity|area|cycle|process'),
    ('ref',        r'^(ref|no\.?|s/?n|#|id|w/?p|wp ref)\b'),
    ('owner',      r'owner|responsib'),
    ('result',     r'result|conclusion|exception|finding|observation|remark|comment'),
    ('freq',       r'frequen|type of control|manual|automated|preventive|detective'),
]
def roles_for(cells):
    out = {}
    for i, c in enumerate(cells):
        c2 = c.strip().lower()
        if not c2 or len(c2) > 60: continue
        for role, rx in ROLES:
            if role in out: continue
            if re.search(rx, c2): out[role] = i; break
    return out

def parse_table(content):
    lines = [l for l in content.split('\n') if l.strip() and not l.startswith('[truncated')]
    rows = [[c.strip() for c in l.split(' | ')] for l in lines]
    # header: first row within the first 15 that has a risk or control column and at least 2 roles
    hdr_i = None
    for i, r in enumerate(rows[:15]):
        ro = roles_for(r)
        if ('risk' in ro or 'control' in ro) and len(ro) >= 2: hdr_i, roles = i, ro; break
    if hdr_i is None: return None, []
    out = []; prev = {}
    for r in rows[hdr_i + 1:]:
        get = lambda k: (r[roles[k]] if k in roles and roles[k] < len(r) else '').strip()
        rec = {k: get(k) for k in ('ref', 'subprocess', 'objective', 'risk', 'control', 'test', 'owner', 'freq', 'result')}
        if not (rec['risk'] or rec['control'] or rec['test']): continue
        if roles_for(r) == roles: continue                       # repeated header
        filled = []
        for k in ('subprocess', 'objective', 'risk', 'control'):    # merged cells: carry the value down when a lower-level cell is filled
            if not rec[k] and prev.get(k) and (rec['test'] or rec['control'] if k != 'control' else rec['test']): rec[k] = prev[k]; filled.append(k)
        rec['filled'] = ','.join(filled); prev = {k: rec[k] for k in ('subprocess', 'objective', 'risk', 'control') if rec[k]} or prev
        out.append(rec)
    return [rows[hdr_i][i] for i in range(len(rows[hdr_i]))], out

types = [t.strip() for t in a.types.split(',')]
sql = f"""select p.id as pid, p.doc_id, p.kind, p.title, p.content, d.name, d.client, d.engagement, d.year, d.process, d.doctype, d.confidence
          from parts p join docs d on d.id = p.doc_id where p.kind in ('table', 'sheet') and d.status = 'done'
          and (d.doctype in ({','.join('?' * len(types))}) {"or (lower(substr(p.content,1,300)) like '%risk%' and lower(substr(p.content,1,300)) like '%control%')" if a.include_hidden else ''})"""
db.execute('drop table if exists rcm_rows')
db.execute('''create table rcm_rows(id integer primary key, doc_id int, part_id int, client text, engagement text, year text, process text, doctype text,
              confidence int, doc_name text, table_title text, ref text, subprocess text, objective text, risk text, control text, test text,
              owner text, freq text, result text, filled text)''')
n_tables = n_parsed = n_rows = 0; per_doc = Counter(); unparsed = Counter(); headers = Counter()
for p in db.execute(sql, types):
    n_tables += 1
    hdr, recs = parse_table(p['content'])
    if hdr is None: unparsed[p['doctype']] += 1; continue
    n_parsed += 1; headers[' | '.join(h[:20] for h in hdr[:8])] += 1
    for r in recs:
        db.execute('insert into rcm_rows(doc_id,part_id,client,engagement,year,process,doctype,confidence,doc_name,table_title,ref,subprocess,objective,risk,control,test,owner,freq,result,filled) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                   (p['doc_id'], p['pid'], p['client'], p['engagement'], p['year'], p['process'], p['doctype'], p['confidence'], p['name'], p['title'],
                    r['ref'][:50], r['subprocess'][:300], r['objective'][:500], r['risk'][:2000], r['control'][:2000], r['test'][:3000], r['owner'][:100], r['freq'][:100], r['result'][:1000], r['filled']))
        n_rows += 1; per_doc[p['doc_id']] += 1
db.commit()
print(f'tables {n_tables:,}  parsed {n_parsed:,}  rows {n_rows:,}  docs with rows {len(per_doc):,}  tables without a recognisable header: {dict(unparsed)}')
print('most common headers:'); [print(f'  {n:>5,}  {h}') for h, n in headers.most_common(12)]

wb = Workbook(); F = Font(name='Arial', size=10); H = Font(name='Arial', size=10, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F4E78')
def sheet(title, hdr, rows, widths):
    ws = wb.create_sheet(title); ws.append(hdr)
    for c in ws[1]: c.font = H; c.fill = HF
    for r in rows: ws.append([str(x)[:2000] if x is not None else '' for x in r])
    for i, w in enumerate(widths, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = F
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
wb.remove(wb.active)
sheet('Rows (first 20000)', ['Id', 'Client', 'Engagement', 'Year', 'Process', 'DocType', 'Doc', 'Table', 'Ref', 'Subprocess', 'Objective', 'Risk', 'Control', 'Test', 'Owner', 'Freq', 'Result', 'Filled'],
      db.execute('select id,client,engagement,year,process,doctype,doc_name,table_title,ref,subprocess,objective,risk,control,test,owner,freq,result,filled from rcm_rows order by id limit 20000'),
      [6, 24, 28, 6, 24, 14, 36, 10, 8, 24, 30, 50, 50, 50, 12, 12, 24, 10])
sheet('Per document', ['DocId', 'Client', 'Engagement', 'Year', 'Doc', 'Rows'], db.execute('select doc_id, client, engagement, year, doc_name, count(*) from rcm_rows group by doc_id order by 6 desc'), [8, 24, 28, 6, 45, 8])
sheet('Headers', ['Header', 'Tables'], headers.most_common(), [120, 8])
wb.save(a.out); print('wrote', a.out)
