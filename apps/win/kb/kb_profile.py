# Purpose: Stage 1 sanity profile of kb.db: coverage per engagement, scanned PDFs without a text twin, RCMs found by content, processes.
#   python kb_profile.py [--db kb.db] [--out kb_profile.xlsx]
import sqlite3, argparse, re, os
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ap = argparse.ArgumentParser(); ap.add_argument('--db', default='kb.db'); ap.add_argument('--out', default='kb_profile.xlsx'); a = ap.parse_args()
db = sqlite3.connect(a.db); db.row_factory = sqlite3.Row
docs = db.execute('select id, name, ext, doctype, confidence, client, engagement, year, yearsource, process, topfolder, status, chars, error from docs').fetchall()
print(f'{len(docs):,} documents')

def stem(name):   # file name without extension, version/sign words and punctuation, for twin matching
    s = os.path.splitext(name)[0].lower()
    s = re.sub(r'\b(signed|scanned|scan|final|draft|copy|v\d+|version\s*\d+|rev\s*\d+)\b', ' ', s)
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()

# ---- coverage per engagement
eng = defaultdict(lambda: {'docs': 0, 'text': 0, 'scanned': 0, 'chars': 0, 'types': Counter(), 'procs': set(), 'years': set()})
for d in docs:
    if not d['client']: continue
    e = eng[(d['client'], d['engagement'])]
    e['docs'] += 1; e['years'].add(d['year'] or '')
    if d['status'] == 'done': e['text'] += 1; e['chars'] += d['chars'] or 0; e['types'][d['doctype']] += 1
    elif d['error'] and 'no text layer' in d['error']: e['scanned'] += 1
    if d['process']: e['procs'].add(d['process'])
TYPES = ['RCM', 'Audit programme', 'Final report', 'Engagement letter', 'Proposal', 'Fee', 'AC paper', 'Planning memo', 'Walkthrough', 'Issues', 'Closing']

# ---- scanned PDFs: does a text twin exist in the same engagement (same stem, or same doctype with text)?
by_eng_stems = defaultdict(set); by_eng_types = defaultdict(set)
for d in docs:
    if d['status'] == 'done' and d['client']:
        by_eng_stems[(d['client'], d['engagement'])].add(stem(d['name'])); by_eng_types[(d['client'], d['engagement'])].add(d['doctype'])
scanned = []
for d in docs:
    if not (d['error'] and 'no text layer' in d['error']): continue
    k = (d['client'], d['engagement']); st = stem(d['name'])
    twin = 'same name' if st in by_eng_stems[k] else ('same type' if d['doctype'] in by_eng_types[k] else 'none')
    pages = re.search(r'in (\d+) pages', d['error']); pages = int(pages.group(1)) if pages else 0
    scanned.append((d['id'], d['client'], d['engagement'], d['year'], d['doctype'], d['confidence'], d['name'], pages, twin))
tw = Counter(s[8] for s in scanned)
print(f"scanned pdfs: {len(scanned):,}  twin same name {tw['same name']:,}  same type {tw['same type']:,}  none {tw['none']:,}  pages without twin {sum(s[7] for s in scanned if s[8]=='none'):,}")

# ---- RCM-like tables in documents not labelled RCM (header row mentions risk and control)
hidden = db.execute("""select p.doc_id, d.name, d.doctype, d.client, d.engagement, d.year, p.kind, p.title, substr(p.content, 1, 200) as head,
                              length(p.content) as len
                       from parts p join docs d on d.id = p.doc_id
                       where p.kind in ('table', 'sheet') and d.doctype not in ('RCM') and length(p.content) > 400
                         and lower(substr(p.content, 1, 300)) like '%risk%' and lower(substr(p.content, 1, 300)) like '%control%'""").fetchall()
seen = set(); hidden_rows = []
for h in hidden:
    if h['doc_id'] in seen: continue
    seen.add(h['doc_id']); hidden_rows.append((h['doc_id'], h['client'], h['engagement'], h['year'], h['doctype'], h['name'], h['kind'], h['title'], h['len'], h['head'].replace('\n', ' / ')[:150]))
print(f'documents with RCM-like tables not labelled RCM: {len(hidden_rows):,}')

# ---- processes
procs = Counter(); rcm_names = Counter()
for d in docs:
    if d['process']: procs[d['process'].strip()] += 1
    if d['doctype'] == 'RCM':
        n = re.sub(r'\.[a-z]+$', '', d['name'].lower()); n = re.sub(r'\b(rcm|risk|control|matrix|and|&|final|draft|v\d+|\d+)\b', ' ', n)
        rcm_names[re.sub(r'[^a-z]+', ' ', n).strip()[:40]] += 1

# ---- workbook
wb = Workbook(); F = Font(name='Arial', size=10); H = Font(name='Arial', size=10, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F4E78')
def sheet(title, hdr, rows, widths):
    ws = wb.create_sheet(title); ws.append(hdr)
    for c in ws[1]: c.font = H; c.fill = HF
    for r in rows: ws.append(list(r))
    for i, w in enumerate(widths, 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = F
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
ws = wb.active; ws.title = 'Summary'
st = Counter((d['doctype'], d['status']) for d in docs)
ws.append(['DocType', 'done', 'error', 'empty', 'chars M']); 
for c in ws[1]: c.font = H; c.fill = HF
for t in TYPES: ws.append([t, st[(t, 'done')], st[(t, 'error')], st[(t, 'empty')], round(sum(d['chars'] or 0 for d in docs if d['doctype'] == t) / 1e6, 1)])
ws.append([]); ws.append(['engagements (client+folder)', len(eng)]); ws.append(['scanned pdfs', len(scanned)]); ws.append(['  with same-name text twin', tw['same name']])
ws.append(['  with same-type text twin', tw['same type']]); ws.append(['  no twin', tw['none']]); ws.append(['  pages to OCR if only no-twin', sum(s[7] for s in scanned if s[8] == 'none')])
ws.append(['docs with RCM-like tables not labelled RCM', len(hidden_rows)])
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = F
ws.column_dimensions['A'].width = 40
sheet('Engagements', ['Client', 'Engagement', 'Years', 'Docs', 'WithText', 'Scanned', 'CharsK', 'Processes'] + TYPES,
      sorted(([c, e, ' '.join(sorted(y for y in v['years'] if y)), v['docs'], v['text'], v['scanned'], round(v['chars'] / 1000), '; '.join(sorted(v['procs']))[:200]] + [v['types'][t] for t in TYPES]) for (c, e), v in eng.items()),
      [28, 34, 12, 7, 8, 8, 8, 50] + [8] * len(TYPES))
sheet('Scanned', ['DocId', 'Client', 'Engagement', 'Year', 'DocType', 'Conf', 'Name', 'Pages', 'TextTwin'], sorted(scanned, key=lambda s: (s[8] != 'none', -s[7])), [8, 28, 34, 7, 16, 6, 50, 7, 11])
sheet('HiddenRCM', ['DocId', 'Client', 'Engagement', 'Year', 'DocType', 'Name', 'Kind', 'Title', 'Chars', 'HeaderRow'], hidden_rows, [8, 28, 34, 7, 16, 45, 7, 20, 8, 80])
sheet('Processes', ['ProcessHint', 'Docs'], procs.most_common(), [60, 8])
sheet('RCMNames', ['NameWords', 'RCMs'], rcm_names.most_common(300), [45, 8])
wb.save(a.out); print('wrote', a.out)
