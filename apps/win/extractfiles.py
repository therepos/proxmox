# Purpose: offline text extractor. Reads the extraction shortlist, pulls text and tables out of each file on the share into one SQLite database with full-text search. No AI, no internet.
# Run with the Python that passed checkenv.py, from the folder holding extraction_shortlist.xlsx.
#
#   python extractfiles.py                                # everything pending in .\extraction_shortlist.xlsx -> .\kb.db, resumable
#   python extractfiles.py --limit 30                     # test run: first 30 pending files
#   python extractfiles.py --workers 4                    # parallel workers, each with its own Word instance (default 3)
#   python extractfiles.py --only RCM,"Audit programme"   # only these DocType values
#   python extractfiles.py --minconf 2                    # skip confidence-1 rows (drafts, weak keyword types)
#   python extractfiles.py --retry-errors                 # re-attempt rows that errored last time
#   python extractfiles.py --redo-crude                   # re-attempt Word files that only got the raw-text fallback (after installing LibreOffice)
#   python extractfiles.py --soffice "C:\Program Files\LibreOffice\program\soffice.exe"   # if LibreOffice is somewhere unusual
#   python extractfiles.py --stats                        # progress by type and status, then exit
#   python extractfiles.py --search "revenue cut-off"     # try the full-text index
#
# Ctrl+C at any time: everything finished so far is committed. Re-run the same command to continue.
# Output: kb.db  tables docs (one row per file, status), parts (text pieces: body/table/sheet/slide/page/email), parts_fts (FTS5).
#         extract.log  start/stop, 5-minute summaries, one line per failed file. Library noise from workers goes to tmp_extract\worker_*.log.
# Old .doc/.rtf/.ppt are first bulk-converted by LibreOffice, 40 files per launch (phase 1), which is far faster than one launch per
# file; anything that fails there goes through Word/PowerPoint, then single-file LibreOffice, then raw text (phase 2).

import sys, os, io, re, json, time, shutil, sqlite3, argparse, subprocess, threading, queue, datetime, traceback

HERE = os.path.dirname(os.path.abspath(__file__))
IS_WIN = os.name == 'nt'
OFFICE_TIMEOUT = 120          # seconds per Word/PowerPoint conversion (file is already local, so longer means a broken file)
FILE_TIMEOUT = 300            # seconds per file for everything else (pdf, excel, email...); a file over this is marked error and skipped
MAX_PART = 2_000_000          # chars per text piece
MAX_ROWS = 20_000             # rows per sheet
MAX_CELL = 4_000              # chars per cell
TEXT_EXT = {'.docx', '.doc', '.rtf', '.xlsx', '.xlsm', '.xls', '.pptx', '.ppt', '.pdf', '.msg', '.txt'}

# ---------------------------------------------------------------- helper process: Word / PowerPoint conversions
def unblock_legacy_word():
    # Trust Center > File Block Settings: 0 = do not block. Only HKCU, so no admin needed; a domain policy may still override.
    try:
        import winreg
        for ver in ('16.0', '15.0', '14.0'):
            try:
                k = winreg.CreateKey(winreg.HKEY_CURRENT_USER, rf'Software\Microsoft\Office\{ver}\Word\Security\FileBlock')
                for v in ('Word2Files', 'Word60Files', 'Word95Files', 'Word2003Files', 'Word2007Files', 'RtfFiles', 'OpenInProtectedView'):
                    winreg.SetValueEx(k, v, 0, winreg.REG_DWORD, 0)
                winreg.CloseKey(k)
            except OSError: pass
    except Exception: pass

def office_worker():
    import win32com.client, pythoncom, ctypes
    pythoncom.CoInitialize()
    unblock_legacy_word()
    word = ppt = None; pids = {}
    def pid_of(hwnd):
        out = ctypes.c_ulong(); ctypes.windll.user32.GetWindowThreadProcessId(int(hwnd), ctypes.byref(out)); return out.value
    def kill_own():
        for p in set(pids.values()):
            subprocess.run(['taskkill', '/F', '/PID', str(p)], capture_output=True)
    def watch_parent():
        # if the worker process that owns this helper dies (Ctrl+C, crash), kill our own Word/PowerPoint so no hidden instances pile up
        k = ctypes.windll.kernel32; h = k.OpenProcess(0x00100000, False, os.getppid())   # SYNCHRONIZE
        if h: k.WaitForSingleObject(h, 0xFFFFFFFF)
        kill_own(); os._exit(0)
    threading.Thread(target=watch_parent, daemon=True).start()
    for line in sys.stdin:
        req = json.loads(line)
        try:
            if req['app'] == 'word':
                if word is None:
                    word = win32com.client.DispatchEx('Word.Application')
                    word.Visible = False; word.DisplayAlerts = 0; word.AutomationSecurity = 3
                    try:
                        o = word.Options
                        o.WarnBeforeSavingPrintingSendingMarkup = False   # the "contains comments and tracked changes, continue?" prompt
                        o.ConfirmConversions = False; o.DoNotPromptForConvert = True; o.SaveInterval = 0
                        o.CheckSpellingAsYouType = False; o.CheckGrammarAsYouType = False
                    except Exception: pass
                    try: d0 = word.Documents.Add(); pids['word'] = pid_of(d0.ActiveWindow.Hwnd); d0.Close(0)
                    except Exception: pass
                d = word.Documents.Open(req['src'], ReadOnly=True, AddToRecentFiles=False, ConfirmConversions=False,
                                        PasswordDocument='__no_password__', Visible=False)
                if 'word' not in pids:
                    try: pids['word'] = pid_of(d.ActiveWindow.Hwnd)
                    except Exception: pass
                d.SaveAs2(req['dst'], FileFormat=12)      # wdFormatXMLDocument
                d.Close(0)
            else:
                if ppt is None:
                    ppt = win32com.client.DispatchEx('PowerPoint.Application')
                    try: pids['ppt'] = pid_of(ppt.HWND)
                    except Exception: pass
                p = ppt.Presentations.Open(req['src'], ReadOnly=True, Untitled=False, WithWindow=False)
                p.SaveAs(req['dst'], 24)                  # ppSaveAsOpenXMLPresentation
                p.Close()
            print(json.dumps({'ok': True, 'pid': pids.get('word') if req['app'] == 'word' else pids.get('ppt')}), flush=True)
        except Exception as e:
            print(json.dumps({'ok': False, 'error': str(e)[:300], 'pid': pids.get('word') if req['app'] == 'word' else pids.get('ppt')}), flush=True)
    for app in (word, ppt):
        try: app.Quit()
        except Exception: pass
    kill_own()

class Office:
    """Talks to the helper process; restarts it and kills the Office app if a conversion exceeds OFFICE_TIMEOUT."""
    def __init__(self): self.p = None; self.q = None; self.pid = None
    def _start(self):
        self.p = subprocess.Popen([sys.executable, os.path.abspath(__file__), '--office-worker'], stdin=subprocess.PIPE,
                                  stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True, bufsize=1)
        self.q = queue.Queue()
        def pump():
            for line in self.p.stdout: self.q.put(line)
            self.q.put(None)
        threading.Thread(target=pump, daemon=True).start()
    def _kill(self):
        try: self.p.kill()
        except Exception: pass
        if self.pid:      # kill only this helper's own Word; never touch the user's Word or other workers'
            subprocess.run(['taskkill', '/F', '/PID', str(self.pid)], capture_output=True)
        self.p = None; self.pid = None
    def convert(self, app, src, dst):
        if not IS_WIN: raise RuntimeError('no Microsoft Office on this system')
        if self.p is None or self.p.poll() is not None: self._start()
        try:
            self.p.stdin.write(json.dumps({'app': app, 'src': src, 'dst': dst}) + '\n'); self.p.stdin.flush()
            line = self.q.get(timeout=OFFICE_TIMEOUT)
        except queue.Empty:
            self._kill(); raise RuntimeError(f'{app} conversion exceeded {OFFICE_TIMEOUT}s, Office killed and restarted')
        except Exception as e:
            self._kill(); raise RuntimeError(f'{app} helper failed: {e}')
        if line is None: self._kill(); raise RuntimeError(f'{app} helper died')
        r = json.loads(line)
        if r.get('pid'): self.pid = r['pid']
        if not r['ok']: raise RuntimeError(r['error'])
    def close(self):
        if self.p:
            try: self.p.stdin.close(); self.p.wait(10)
            except Exception: self._kill()

def find_soffice(explicit=''):
    if explicit: return explicit if os.path.exists(explicit) else None
    p = shutil.which('soffice') or shutil.which('soffice.exe')
    if p: return p
    import glob
    for pat in (r'C:\Program Files*\LibreOffice*\program\soffice.exe', os.path.join(HERE, 'LibreOffice*', 'App', 'libreoffice', 'program', 'soffice.exe'),
                os.path.join(HERE, '*', 'program', 'soffice.exe')):
        m = glob.glob(pat)
        if m: return m[0]
    return None

def soffice_convert(soffice, src, fmt, tmpdir):
    # headless LibreOffice, own profile per worker so several can run at once; returns the converted file path
    prof = os.path.join(tmpdir, f'lo_profile_{os.getpid()}'); outdir = os.path.join(tmpdir, f'lo_out_{os.getpid()}')
    os.makedirs(outdir, exist_ok=True)
    url = 'file:///' + os.path.abspath(prof).replace('\\', '/')
    cmd = [soffice, f'-env:UserInstallation={url}', '--headless', '--norestore', '--nolockcheck', '--convert-to', fmt, '--outdir', outdir, src]
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=OFFICE_TIMEOUT)
    except subprocess.TimeoutExpired as e:
        raise RuntimeError(f'libreoffice exceeded {OFFICE_TIMEOUT}s')
    out = os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + '.' + fmt)
    if not os.path.exists(out): raise RuntimeError('libreoffice: ' + ((r.stderr or r.stdout or '').strip()[-200:] or 'no output file'))
    return out

# ---------------------------------------------------------------- text handlers: each returns [(kind, title, content)]
def clean(s):
    if s is None: return ''
    s = str(s).replace('\x00', '').replace('\r\n', '\n').replace('\r', '\n')
    return re.sub(r'[ \t\xa0]+', ' ', s).strip()

def cellstr(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    if isinstance(v, (datetime.datetime, datetime.date)): return v.isoformat()[:10]
    s = clean(v)
    return s[:MAX_CELL]

def rows_to_text(rows):
    out = []; n = 0
    for r in rows:
        cells = [cellstr(c) for c in r]
        while cells and cells[-1] == '': cells.pop()
        if not any(cells): continue
        out.append(' | '.join(cells)); n += 1
        if n >= MAX_ROWS: out.append(f'[truncated at {MAX_ROWS} rows]'); break
    return '\n'.join(out)

def h_docx(path):
    import docx
    from docx.table import _Cell
    d = docx.Document(path)
    parts = []
    body = '\n'.join(clean(p.text) for p in d.paragraphs if clean(p.text))
    if body: parts.append(('body', '', body))
    for i, t in enumerate(d.tables, 1):
        # read cells straight from the XML: python-docx's row.cells is strict about grid layout and fails on converted files
        rows = [[_Cell(tc, t).text for tc in tr.tc_lst] for tr in t._tbl.tr_lst]
        txt = rows_to_text(rows)
        if txt: parts.append(('table', f'table {i}', txt))
    return parts

def h_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        txt = rows_to_text(ws.iter_rows(values_only=True))
        if txt: parts.append(('sheet', ws.title, txt))
    wb.close()
    return parts

def h_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path, on_demand=True, formatting_info=False)
    parts = []
    for name in wb.sheet_names():
        ws = wb.sheet_by_name(name)
        rows = ([xlrd.xldate_as_datetime(c.value, wb.datemode) if c.ctype == xlrd.XL_CELL_DATE and c.value else c.value
                 for c in ws.row(r)] for r in range(ws.nrows))
        txt = rows_to_text(rows)
        if txt: parts.append(('sheet', name, txt))
        wb.unload_sheet(name)
    return parts

def h_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, s in enumerate(prs.slides, 1):
        lines = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = clean(sh.text_frame.text)
                if t: lines.append(t)
            if getattr(sh, 'has_table', False) and sh.has_table:
                lines.append(rows_to_text([[c.text for c in row.cells] for row in sh.table.rows]))
        if s.has_notes_slide:
            t = clean(s.notes_slide.notes_text_frame.text)
            if t: lines.append('[notes] ' + t)
        txt = '\n'.join(l for l in lines if l)
        if txt: parts.append(('slide', f'slide {i}', txt))
    return parts

def h_pdf(path):
    from pypdf import PdfReader
    r = PdfReader(path, strict=False)
    if r.is_encrypted:
        try: r.decrypt('')
        except Exception: raise RuntimeError('pdf is password protected')
    parts = []
    for i, pg in enumerate(r.pages, 1):
        try:
            res = pg.get('/Resources')
            if res is not None and '/Font' not in res and '/XObject' not in res: continue   # image-only or blank page, no text possible
            t = clean(pg.extract_text() or '')
        except Exception as e: t = f'[page {i}: extract failed {e}]'
        if t: parts.append(('page', f'page {i}', t))
    if not parts and len(r.pages): raise RuntimeError(f'no text layer in {len(r.pages)} pages (scanned image pdf, needs OCR)')
    return parts

def crude_text(raw):
    out = []
    for m in re.finditer(rb'(?:[\x20-\x7e\x07\x0b\x0d\xa0-\xff]){12,}', raw):            # 8-bit text pieces
        out.append(m.group().decode('cp1252', 'replace'))
    for m in re.finditer(rb'(?:[\x20-\x7e\x07\x0b\x0d]\x00){12,}', raw):                    # UTF-16 pieces
        out.append(m.group().decode('utf-16le', 'replace'))
    out = [t for t in out if sum(c.isascii() and (c.isalnum() or c in ' .,;:()-/%$|\x07\x0b\r') for c in t) / len(t) >= 0.85]   # drop binary noise
    txt = '\n'.join(out).replace('\x07\x07', '\n').replace('\x07', ' | ').replace('\x0b', '\n').replace('\r', '\n')
    return re.sub(r'\n\s*\n+', '\n', clean(txt))

def h_doc_crude(path):
    # Fallback when Word cannot open the file (policy block, timeout, corruption): pull readable runs from the
    # WordDocument stream. Works for Word 6/95/97-2003. Tables flatten to 'cell | cell' lines. Lower fidelity than Word.
    import olefile
    if not olefile.isOleFile(path): raise RuntimeError('not an OLE file')
    with olefile.OleFileIO(path) as ole:
        if not ole.exists('WordDocument'): raise RuntimeError('no WordDocument stream')
        raw = ole.openstream('WordDocument').read()
    txt = crude_text(raw)
    if len(re.sub(r'[^A-Za-z]', '', txt)) < 40: raise RuntimeError('crude read found no text')
    return [('body-crude', 'raw text, Word could not open the file', txt)]

def h_txt(path):
    with open(path, 'rb') as fh: b = fh.read()
    for enc in ('utf-8-sig', 'utf-16', 'cp1252'):
        try: t = b.decode(enc); break
        except Exception: t = b.decode('latin-1')
    t = clean(t)
    return [('body', '', t)] if t else []

def h_msg(path, tmpdir, office, soffice=None, depth=0):
    import extract_msg
    m = extract_msg.openMsg(path)
    parts = []
    head = '\n'.join(f'{k}: {clean(v)}' for k, v in (('Subject', m.subject), ('From', m.sender), ('To', m.to), ('Cc', m.cc), ('Date', m.date)) if v)
    parts.append(('email', 'header', head))
    body = clean(m.body or '')
    if body: parts.append(('email', 'body', body))
    for a in m.attachments:
        name = getattr(a, 'longFilename', None) or getattr(a, 'shortFilename', None) or ''
        ext = os.path.splitext(name)[1].lower()
        if depth or ext not in TEXT_EXT or ext == '.msg' or not isinstance(getattr(a, 'data', None), (bytes, bytearray)):
            if name: parts.append(('attachment', name, '[attachment not extracted]'))
            continue
        ap = os.path.join(tmpdir, 'att_' + re.sub(r'[^\w.]+', '_', name)[-80:])
        with open(ap, 'wb') as fh: fh.write(a.data)
        try:
            for kind, title, content in extract(ap, ext, tmpdir, office, soffice):
                parts.append(('attachment', f'{name} / {title}' if title else name, content))
        except Exception as e:
            parts.append(('attachment', name, f'[attachment failed: {e}]'))
        finally:
            try: os.remove(ap)
            except Exception: pass
    m.close()
    return parts

def sniff(local, ext):
    # files are often saved with the wrong extension; route by content
    try:
        with open(local, 'rb') as fh: head = fh.read(8)
    except Exception: return ext
    if head.startswith(b'%PDF'): return '.pdf'
    if head.startswith(b'PK\x03\x04'):
        import zipfile
        try:
            names = zipfile.ZipFile(local).namelist()
            if any(n.startswith('word/') for n in names): return '.docx'
            if any(n.startswith('xl/') for n in names): return '.xlsx'
            if any(n.startswith('ppt/') for n in names): return '.pptx'
        except Exception: pass
        return ext
    if head.startswith(b'\xd0\xcf\x11\xe0'):
        import olefile
        try:
            with olefile.OleFileIO(local) as ole:
                if ole.exists('WordDocument'): return '.doc'
                if ole.exists('Workbook') or ole.exists('Book'): return '.xls'
                if ole.exists('PowerPoint Document'): return '.ppt'
                if ole.exists('__properties_version1.0'): return '.msg'
        except Exception: pass
    if head.startswith(b'{\\rtf'): return '.rtf'
    return ext

def extract(local, ext, tmpdir, office, soffice=None, did=None):
    ext = sniff(local, ext)
    if did is not None and ext in ('.doc', '.rtf', '.ppt'):       # phase-1 bulk conversion already done?
        cached = os.path.join(tmpdir, 'converted', f'{did}.docx' if ext != '.ppt' else f'{did}.pptx')
        if os.path.exists(cached):
            try: return h_docx(cached) if ext != '.ppt' else h_pptx(cached)
            except Exception as e: cache_err = f'cached {os.path.basename(cached)} unreadable: {e}'
        else: cache_err = None
    else: cache_err = None
    if ext == '.docx': return h_docx(local)
    if ext in ('.doc', '.rtf'):
        errs = []
        dst = local + '.docx'
        try:
            office.convert('word', local, dst); parts = h_docx(dst)
            if cache_err: parts.append(('note', 'conversion', cache_err))
            return parts
        except Exception as e: errs.append(f'word: {e}')
        finally:
            if os.path.exists(dst): os.remove(dst)
        if soffice:
            try:
                out = soffice_convert(soffice, local, 'docx', tmpdir)
                try: return h_docx(out)
                finally: os.remove(out)
            except Exception as e: errs.append(f'libreoffice: {e}')
        if ext == '.doc':
            try: parts = h_doc_crude(local)
            except Exception as e: errs.append(f'raw: {e}'); raise RuntimeError(' | '.join(errs))
            parts.append(('note', 'conversion', 'raw text used, converters failed: ' + ' | '.join(errs)))
            return parts
        raise RuntimeError(' | '.join(errs))
    if ext in ('.xlsx', '.xlsm'): return h_xlsx(local)
    if ext == '.xls':
        try: return h_xls(local)
        except Exception as e:
            # some ".xls" are really xlsx, html or csv
            try: return h_xlsx(local)
            except Exception: pass
            dst = local + '.docx'
            try: office.convert('word', local, dst); return h_docx(dst)
            except Exception: raise e
            finally:
                if os.path.exists(dst): os.remove(dst)
    if ext == '.pptx': return h_pptx(local)
    if ext == '.ppt':
        errs = []
        dst = local + '.pptx'
        try:
            office.convert('ppt', local, dst); return h_pptx(dst)
        except Exception as e: errs.append(f'powerpoint: {e}')
        finally:
            if os.path.exists(dst): os.remove(dst)
        if soffice:
            try:
                out = soffice_convert(soffice, local, 'pptx', tmpdir)
                try: return h_pptx(out)
                finally: os.remove(out)
            except Exception as e: errs.append(f'libreoffice: {e}')
        raise RuntimeError(' | '.join(errs))
    if ext == '.pdf': return h_pdf(local)
    if ext == '.msg': return h_msg(local, tmpdir, office, soffice)
    if ext == '.txt': return h_txt(local)
    raise RuntimeError(f'unsupported extension {ext}')

# ---------------------------------------------------------------- database
def open_db(dbpath):
    db = sqlite3.connect(dbpath, timeout=60)
    db.executescript('''
    PRAGMA journal_mode=WAL;
    CREATE TABLE IF NOT EXISTS docs(id INTEGER PRIMARY KEY, path TEXT UNIQUE, name TEXT, ext TEXT, size INTEGER, modified TEXT,
        doctype TEXT, confidence INTEGER, client TEXT, engagement TEXT, year TEXT, yearsource TEXT, process TEXT, topfolder TEXT, duplicates INTEGER,
        status TEXT DEFAULT 'pending', parts INTEGER, chars INTEGER, error TEXT, seconds REAL, done_at TEXT, copy_s REAL);
    CREATE INDEX IF NOT EXISTS ix_docs_status ON docs(status);
    CREATE TABLE IF NOT EXISTS parts(id INTEGER PRIMARY KEY, doc_id INTEGER, seq INTEGER, kind TEXT, title TEXT, content TEXT);
    CREATE INDEX IF NOT EXISTS ix_parts_doc ON parts(doc_id);
    CREATE VIRTUAL TABLE IF NOT EXISTS parts_fts USING fts5(title, content, content='parts', content_rowid='id', tokenize='porter unicode61');
    ''')
    try: db.execute('alter table docs add column copy_s REAL')
    except sqlite3.OperationalError: pass
    return db

def load_list(db, xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb['Shortlist'] if 'Shortlist' in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True); hdr = [str(h) for h in next(rows)]
    ix = {h: i for i, h in enumerate(hdr)}
    def g(r, k): v = r[ix[k]] if k in ix else None; return v
    n = 0
    for r in rows:
        if not r or not g(r, 'FullPath'): continue
        cur = db.execute('''INSERT OR IGNORE INTO docs(path,name,ext,size,modified,doctype,confidence,client,engagement,year,yearsource,process,topfolder,duplicates)
                            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                         (g(r, 'FullPath'), g(r, 'FileName'), (g(r, 'Ext') or '').lower(), g(r, 'SizeBytes'), str(g(r, 'Modified') or ''), g(r, 'DocType'),
                          g(r, 'Confidence'), g(r, 'Client'), g(r, 'Engagement'), str(g(r, 'Year') or ''), g(r, 'YearSource'), g(r, 'ProcessHint'), g(r, 'TopFolder'), g(r, 'Duplicates')))
        n += cur.rowcount
    db.commit(); wb.close()
    return n

def longpath(p):
    if not IS_WIN or p.startswith('\\\\?\\'): return p
    return '\\\\?\\UNC\\' + p[2:] if p.startswith('\\\\') else '\\\\?\\' + p

LOG = os.path.join(HERE, 'extract.log')
def log(msg):
    try:
        with open(LOG, 'a', encoding='utf-8') as fh: fh.write(f'{datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")}  {msg}\n')
    except Exception: pass

def stats(db):
    print('\nstatus by document type')
    print(f"{'DocType':20}{'pending':>9}{'done':>9}{'empty':>9}{'error':>9}{'chars M':>10}")
    for row in db.execute('''select doctype, sum(status='pending'), sum(status='done'), sum(status='empty'), sum(status='error'), sum(coalesce(chars,0))/1e6
                             from docs group by doctype order by doctype'''):
        print(f"{row[0]:20}{row[1]:>9,}{row[2]:>9,}{row[3]:>9,}{row[4]:>9,}{row[5]:>10.1f}")
    row = db.execute("select count(*), sum(status='done'), sum(status='error'), sum(status='empty') from docs").fetchone()
    print(f"total {row[0]:,}   done {row[1]:,}   error {row[2]:,}   empty {row[3]:,}")
    print('\ntiming by extension (processed files)')
    print(f"{'ext':8}{'files':>8}{'avg copy s':>12}{'avg total s':>13}{'max s':>8}{'crude':>7}")
    for row in db.execute('''select ext, count(*), avg(copy_s), avg(seconds), max(seconds),
                             sum(exists(select 1 from parts p where p.doc_id=docs.id and p.kind='body-crude'))
                             from docs where status!='pending' group by ext order by 2 desc'''):
        print(f"{row[0]:8}{row[1]:>8,}{(row[2] or 0):>12.1f}{(row[3] or 0):>13.1f}{(row[4] or 0):>8.0f}{row[5]:>7}")
    print('\ntop errors')
    for e, c in db.execute("select substr(error,1,90), count(*) from docs where status='error' group by 1 order by 2 desc limit 12"): print(f"  {c:>6,}  {e}")

def search(db, q):
    if '"' not in q: q = ' '.join('"' + t.replace('"', '') + '"' for t in q.split())   # plain words: AND of exact tokens, hyphens safe
    for name, doctype, client, title, snip in db.execute('''select d.name, d.doctype, d.client, p.title, snippet(parts_fts, 1, '[', ']', ' ... ', 18)
        from parts_fts join parts p on p.id = parts_fts.rowid join docs d on d.id = p.doc_id where parts_fts match ? order by rank limit 15''', (q,)):
        print(f"{doctype:16} {(client or '')[:22]:22} {name[:45]:45} {title[:14]:14} {snip}")

def hms(s): return str(datetime.timedelta(seconds=int(s))) if s >= 0 else '--:--:--'

# ---------------------------------------------------------------- phase 1: bulk LibreOffice conversion, 40 files per launch
def w_bulk(chunk):
    # chunk: [(did, path, ext)] all .doc/.rtf or all .ppt. Copies to a batch folder, one soffice call, moves results to converted/<id>.<fmt>
    fmt = 'pptx' if chunk[0][2] == '.ppt' else 'docx'
    tmp = W['tmp']; bdir = os.path.join(tmp, f'batch_{os.getpid()}'); conv = os.path.join(tmp, 'converted')
    shutil.rmtree(bdir, ignore_errors=True); os.makedirs(bdir); os.makedirs(conv, exist_ok=True)
    srcs = []
    for did, path, ext in chunk:
        local = os.path.join(bdir, f'{did}{ext}')
        try: shutil.copyfile(longpath(path), local); srcs.append(local)
        except Exception: pass
    try: W['busy'][os.getpid()] = (f'bulk convert {len(srcs)} {fmt} files', time.time())
    except Exception: pass
    ok = 0
    if srcs:
        prof = os.path.join(tmp, f'lo_profile_{os.getpid()}'); url = 'file:///' + os.path.abspath(prof).replace('\\', '/')
        cmd = [W['soffice'], f'-env:UserInstallation={url}', '--headless', '--norestore', '--nolockcheck', '--convert-to', fmt, '--outdir', bdir] + srcs
        try: subprocess.run(cmd, capture_output=True, timeout=OFFICE_TIMEOUT * 4)
        except subprocess.TimeoutExpired: pass
        except Exception: pass
        for did, path, ext in chunk:
            out = os.path.join(bdir, f'{did}.{fmt}')
            if os.path.exists(out) and os.path.getsize(out) > 0:
                try: os.replace(out, os.path.join(conv, f'{did}.{fmt}')); ok += 1
                except Exception: pass
    shutil.rmtree(bdir, ignore_errors=True)
    try: W['busy'][os.getpid()] = ('', 0)
    except Exception: pass
    return len(chunk), ok

# ---------------------------------------------------------------- worker process: copy + extract one file, return parts
W = {}
def w_init(tmpdir, soffice, busy_map):
    W['tmp'] = tmpdir; W['office'] = Office(); W['soffice'] = soffice; W['busy'] = busy_map
    import warnings, logging
    warnings.filterwarnings('ignore')
    for name in ('pypdf', 'openpyxl', 'extract_msg', 'RTFDE'): logging.getLogger(name).setLevel(logging.CRITICAL)
    try: sys.stderr = open(os.path.join(tmpdir, f'worker_{os.getpid()}.log'), 'a', encoding='utf-8', errors='replace')
    except Exception: pass
    import signal, atexit; signal.signal(signal.SIGINT, signal.SIG_IGN)     # parent handles Ctrl+C
    atexit.register(W['office'].close)

def w_one(item):
    did, path, ext, name, doctype = item
    ts = time.time(); local = os.path.join(W['tmp'], f'{did}_{os.getpid()}{ext}')
    try: W['busy'][os.getpid()] = (name, ts)
    except Exception: pass
    status, err, parts, copy_s = 'done', None, [], None
    try:
        try: shutil.copyfile(longpath(path), local); copy_s = round(time.time() - ts, 2)
        except Exception as e: raise RuntimeError(f'copy: {e}')
        box = {}
        def run():
            try: box['parts'] = extract(local, ext, W['tmp'], W['office'], W['soffice'], did)
            except BaseException as e: box['err'] = e
        th = threading.Thread(target=run, daemon=True); th.start(); th.join(FILE_TIMEOUT)
        if th.is_alive():
            # raise TimeoutError inside the runaway reader thread; it takes effect as soon as that thread is back in Python code
            import ctypes
            ctypes.pythonapi.PyThreadState_SetAsyncExc(ctypes.c_ulong(th.ident), ctypes.py_object(TimeoutError))
            th.join(10)
            raise RuntimeError(f'timeout: reader still busy after {FILE_TIMEOUT}s, file skipped')
        if 'err' in box: raise box['err']
        parts = box['parts']
        parts = [(k, t, c[:MAX_PART]) for k, t, c in parts if c]
        if not parts: status, err = 'empty', 'no text found'
    except Exception as e:
        status, err = 'error', (str(e) or e.__class__.__name__)[:400]
    finally:
        if os.path.exists(local):
            try: os.remove(local)
            except Exception: pass
    try: W['busy'][os.getpid()] = ('', 0)
    except Exception: pass
    return did, name, status, err, parts, round(time.time() - ts, 2), copy_s

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--list', default=os.path.join(HERE, 'extraction_shortlist.xlsx'))
    ap.add_argument('--db', default=os.path.join(HERE, 'kb.db'))
    ap.add_argument('--only', default='', help='comma separated DocType values')
    ap.add_argument('--minconf', type=int, default=1)
    ap.add_argument('--limit', type=int, default=0)
    ap.add_argument('--workers', type=int, default=3)
    ap.add_argument('--retry-errors', action='store_true')
    ap.add_argument('--redo-crude', action='store_true')
    ap.add_argument('--soffice', default='')
    ap.add_argument('--no-bulk', action='store_true', help='skip the phase-1 bulk LibreOffice conversion')
    ap.add_argument('--stats', action='store_true')
    ap.add_argument('--search', default='')
    ap.add_argument('--office-worker', action='store_true')
    a = ap.parse_args()
    if a.office_worker: return office_worker()

    db = open_db(a.db)
    if a.search: return search(db, a.search)
    if a.stats: return stats(db)          # read-only, safe while another instance is running
    if os.path.exists(a.list):
        n = load_list(db, a.list)
        if n: print(f'loaded {n:,} new rows from {os.path.basename(a.list)}')
    if a.retry_errors: db.execute("update docs set status='pending', error=NULL where status='error'"); db.commit()
    if a.redo_crude:
        db.execute("update docs set status='pending' where id in (select doc_id from parts where kind='body-crude')"); db.commit()
    sql = "select id, path, ext, name, doctype from docs where status='pending' and confidence >= ?"; args = [a.minconf]
    if a.only:
        kinds = [k.strip() for k in a.only.split(',') if k.strip()]
        sql += f" and doctype in ({','.join('?' * len(kinds))})"; args += kinds
    sql += ' order by id'
    if a.limit: sql += f' limit {int(a.limit)}'
    todo = db.execute(sql, args).fetchall()
    total = db.execute('select count(*) from docs').fetchone()[0]
    done0 = db.execute("select count(*) from docs where status!='pending'").fetchone()[0]
    print(f'db: {a.db}\n{total:,} rows in list, {done0:,} already processed, {len(todo):,} to do now')
    if not todo: return stats(db)

    first = todo[0][1]
    root = first[:first.find('\\', first.find('\\', 2) + 1)] if first.startswith('\\\\') else os.path.dirname(first)
    if not os.path.isdir(longpath(root)):
        print(f'share not reachable: {root}\nopen it in Explorer (it may ask for a login), then re-run.'); return
    tmpdir = os.path.join(HERE, 'tmp_extract'); os.makedirs(tmpdir, exist_ok=True)
    import multiprocessing as mp
    soffice = find_soffice(a.soffice)
    print('libreoffice:', soffice or 'not found (Word 95 files get raw text only)')
    mgr = mp.Manager(); busy_map = mgr.dict()
    pool = mp.Pool(max(1, a.workers), initializer=w_init, initargs=(tmpdir, soffice, busy_map), maxtasksperchild=300)
    log(f'start: {len(todo):,} to do, workers {a.workers}, libreoffice {soffice}')

    # phase 1: bulk-convert legacy Office files not yet in the converted cache
    conv = os.path.join(tmpdir, 'converted'); os.makedirs(conv, exist_ok=True)
    if soffice and not a.no_bulk:
        need = [(d, p, e) for d, p, e, _, _ in todo if e in ('.doc', '.rtf', '.ppt') and not os.path.exists(os.path.join(conv, f'{d}.' + ('pptx' if e == '.ppt' else 'docx')))]
        docs = [x for x in need if x[2] != '.ppt']; ppts = [x for x in need if x[2] == '.ppt']
        chunks = [docs[i:i + 40] for i in range(0, len(docs), 40)] + [ppts[i:i + 40] for i in range(0, len(ppts), 40)]
        if chunks:
            print(f'phase 1: bulk converting {len(need):,} old Word/PowerPoint files with LibreOffice, {len(chunks)} batches')
            t1 = time.time(); done1 = 0; ok1 = 0
            try:
                for k, (cnt, ok) in enumerate(pool.imap_unordered(w_bulk, chunks), 1):
                    done1 += cnt; ok1 += ok; el = time.time() - t1
                    print(f'\rbatch {k}/{len(chunks)}  {done1:,} files, {ok1:,} converted  {hms(el)}  eta {hms(el / k * (len(chunks) - k))}'.ljust(110), end='', flush=True)
            except KeyboardInterrupt:
                print('\nstopped during bulk conversion. converted files are kept; re-run to continue.'); pool.terminate(); return
            print(f'\nphase 1 done: {ok1:,} of {done1:,} converted in {hms(time.time() - t1)}; the rest go through Word/LibreOffice one by one')
            log(f'phase 1: {ok1:,} of {done1:,} bulk-converted in {hms(time.time() - t1)}')

    print(f'phase 2: extracting text from {len(todo):,} files')
    t0 = time.time(); n = 0; nerr = 0; nempty = 0; copyfail = 0
    total_todo = len(todo); stop_ui = threading.Event(); hist = []   # (time, n) samples for a recent-rate ETA

    def ui():   # heartbeat: refreshes every second even while workers are busy, permanent summary line every 5 min
        lastsum = time.time()
        while not stop_ui.wait(1):
            now = time.time(); el = now - t0
            hist.append((now, n)); del hist[:-600]
            old = next((h for h in hist if now - h[0] >= 300), hist[0])
            rate = (n - old[1]) / max(now - old[0], 1) * 60 if now - old[0] > 30 else n / max(el, 1) * 60
            eta = (total_todo - n) / rate * 60 if rate > 0 else -1
            busy = sorted(((now - t, nm) for nm, t in busy_map.values() if nm), reverse=True)
            cur = f'{len(busy)} busy, longest {busy[0][0]:.0f}s: {busy[0][1]}' if busy else 'waiting for workers'
            line = f'{n:,}/{total_todo:,}  err {nerr:,}  {rate:4.1f}/min  {hms(el)}  eta {hms(eta)}  | {cur}'
            width = max(60, shutil.get_terminal_size((120, 20)).columns - 1)
            print('\r' + line[:width].ljust(width), end='', flush=True)
            if now - lastsum >= 300:
                lastsum = now
                msg = f'{n:,} of {total_todo:,} done  err {nerr:,}  empty {nempty:,}  {rate:4.1f}/min  eta {hms(eta)}'
                print(f'\r{datetime.datetime.now().strftime("%H:%M")}  {msg}'.ljust(width)); log(msg)
    threading.Thread(target=ui, daemon=True).start()
    try:
        for did, name, status, err, parts, secs, copy_s in pool.imap_unordered(w_one, todo, chunksize=1):
            n += 1
            if status == 'error': nerr += 1; log(f'error  {name}  |  {err}')
            copyfail = copyfail + 1 if (err or '').startswith('copy:') else 0
            if copyfail >= 30:
                print(f'\n{copyfail} files in a row could not be copied from the share, it is probably disconnected. stopping.')
                print('reconnect the share, then re-run with --retry-errors to redo the failed rows.')
                raise KeyboardInterrupt
            if status == 'empty': nempty += 1
            for pid, t, c in db.execute('select id, title, content from parts where doc_id=?', (did,)).fetchall():   # retry: drop old text
                db.execute("insert into parts_fts(parts_fts, rowid, title, content) values('delete', ?, ?, ?)", (pid, t, c))
            db.execute('delete from parts where doc_id=?', (did,))
            for i, (k, t, c) in enumerate(parts, 1):
                cur = db.execute('insert into parts(doc_id,seq,kind,title,content) values(?,?,?,?,?)', (did, i, k, t, c))
                db.execute('insert into parts_fts(rowid,title,content) values(?,?,?)', (cur.lastrowid, t, c))
            db.execute('update docs set status=?, parts=?, chars=?, error=?, seconds=?, copy_s=?, done_at=? where id=?',
                       (status, len(parts), sum(len(c) for _, _, c in parts), err, secs, copy_s, datetime.datetime.now().strftime('%Y-%m-%d %H:%M'), did))
            if n % 25 == 0: db.commit()
        pool.close(); pool.join()
    except KeyboardInterrupt:
        print('\nstopped. re-run the same command to continue.'); log(f'stopped by user after {n:,} files')
        pool.terminate()      # helpers lose their pipe and quit their own Word/PowerPoint; the user's Office is never touched
    finally:
        stop_ui.set(); db.commit()
    print(f'\nfinished {n:,} files in {hms(time.time() - t0)}'); log(f'finished {n:,} files, err {nerr:,}, empty {nempty:,}, in {hms(time.time() - t0)}')
    stats(db)

if __name__ == '__main__':
    main()
